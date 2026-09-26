"""Тёмный шаблон (VK WorkSpace): контраст по фактическому фону (ADR-034).

Регресс: макеты задают фон свойством `p:bg = #000000`, а рендер и аудит брали
фон из светлой палитры темы — контент выходил чёрным на чёрном и «пропадал»,
хотя audit_matrix рапортовал «чисто». Здесь проверяется единый алгоритм:
заливка фигуры → декор макета → фон p:bg.
"""
from __future__ import annotations

import io
import re

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

from app.audit.checks import Audit, _contrast, _hex
from app.layout.engine import DesignContext
from app.models.deck import Block, Deck, Slide, SlideType
from app.render.images import hex_to_rgb
from app.render.pptx_renderer import Renderer

ZERO_TEXT = re.compile(r"^[\s0/]+$")


def _pitch_like_deck() -> Deck:
    return Deck(title="Защита проекта: цифровой дизайнер", slides=[
        Slide(slide_type=SlideType.TITLE, heading="Цифровой дизайнер презентаций",
              subheading="Проект для хакатона"),
        Slide(slide_type=SlideType.SECTION, heading="Проблема"),
        Slide(slide_type=SlideType.CONTENT, heading="Проблема: разрыв между смыслом и дизайном",
              blocks=[Block(kind="bullets", items=[
                  "Каждую неделю рождаются десятки презентаций.",
                  "Авторы тратят часы на ручную переверстку.",
                  "Модели нарушают дизайн-систему шаблона.",
              ])]),
        Slide(slide_type=SlideType.CONTENT, heading="Архитектура: пайплайн из восьми слоёв",
              blocks=[Block(kind="steps", items=[
                  "Парсинг шаблона в JSON-профиль.",
                  "Планирование содержания.",
                  "Вёрстка трёх вариантов.",
                  "Аудит и авто-фиксы.",
              ])]),
        Slide(slide_type=SlideType.CONTENT, heading="Результаты: 180 тестов и 24 комбинации",
              blocks=[Block(kind="factoids", factoids=[
                  {"value": "180", "label": "автоматических тестов"},
                  {"value": "24/24", "label": "комбинации без ошибок"},
                  {"value": "34–76 сек", "label": "время генерации"},
                  {"value": "5 мин", "label": "бюджет задания"},
              ])]),
        Slide(slide_type=SlideType.FINAL, heading="Спасибо за внимание"),
    ])


def _threshold(size: float, bold: bool = False) -> float:
    return 3.0 if (size >= 18 or (size >= 14 and bold)) else 4.5


def test_dark_template_text_contrast(profile_of, template_workspace):
    """Все тексты на тёмном шаблоне читаются на фактическом фоне."""
    profile = profile_of(template_workspace)
    dc = DesignContext.from_profile(profile)
    deck = _pitch_like_deck()
    for variant in ("compact", "cards", "split"):
        renderer = Renderer(profile, variant=variant, template_bytes=template_workspace)
        pptx = renderer.render_deck(deck, dc)
        audit = Audit(profile).audit(deck, pptx)
        codes = {issue["code"] for issue in audit["issues"]}
        assert "text_invisible" not in codes, audit["issues"][:5]
        assert "contrast_too_low" not in codes, audit["issues"][:5]
        assert audit["errors"] == 0, audit["issues"][:5]

        prs = Presentation(io.BytesIO(pptx))
        auditor = Audit(profile)
        auditor.W, auditor.H = prs.slide_width, prs.slide_height
        auditor._deck = deck
        checked = 0
        for slide in prs.slides:
            for shape in slide.shapes:
                if not shape.has_text_frame or not shape.text_frame.text.strip():
                    continue
                bg = auditor._effective_bg(slide, shape)
                for p in shape.text_frame.paragraphs:
                    for run in p.runs:
                        if not run.text.strip():
                            continue
                        try:
                            color = "#" + str(run.font.color.rgb)
                        except Exception:  # noqa: BLE001 — без явного цвета
                            continue
                        size = run.font.size.pt if run.font.size else 14.0
                        bold = bool(run.font.bold)
                        ratio = _contrast(*hex_to_rgb(color), *_hex(bg))
                        assert ratio >= _threshold(size, bold), (
                            f"{variant}: «{run.text[:30]}» {color} на {bg} "
                            f"контраст {ratio:.2f} при кегле {size}")
                        checked += 1
        assert checked >= 10


def test_title_placeholders_cleared(profile_of, template_workspace):
    """Фоновый шум «0 / 0 / 0» и пустые плейсхолдеры не попадают в вывод."""
    profile = profile_of(template_workspace)
    dc = DesignContext.from_profile(profile)
    deck = _pitch_like_deck()
    renderer = Renderer(profile, variant="compact", template_bytes=template_workspace)
    pptx = renderer.render_deck(deck, dc)
    prs = Presentation(io.BytesIO(pptx))

    title_slide = list(prs.slides)[0]
    for shape in title_slide.shapes:
        if shape.has_text_frame:
            for line in shape.text_frame.text.splitlines():
                assert not ZERO_TEXT.match(line), f"текст-заглушка «{line}»"

    slide_w = prs.slide_width / 914400
    slide_h = prs.slide_height / 914400
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            for shape in _pictures_deep(layout.shapes):
                width = (shape.width or 0) / 914400
                height = (shape.height or 0) / 914400
                if width * height < 0.25 * slide_w * slide_h:
                    continue
                brightness = _brightness(shape.image.blob)
                if brightness >= 0.12:
                    continue
                x = (shape.left or 0) / 914400
                y = (shape.top or 0) / 914400
                assert (x >= -0.1 and y >= -0.1 and x + width <= slide_w + 0.1
                        and y + height <= slide_h + 0.1), (
                    f"тёмная картинка-шум осталась: {layout.name}/{shape.name}")


def test_contrast_check_matches_render(profile_of, template_workspace):
    """Аудит и рендер определяют фон одним и тем же алгоритмом."""
    profile = profile_of(template_workspace)
    renderer = Renderer(profile, variant="compact", template_bytes=template_workspace)
    auditor = Audit(profile)
    content_layouts = [l for l in profile["layouts"] if l.get("background")]
    assert content_layouts, "у VK WorkSpace фон p:bg должен читаться"
    for layout in content_layouts:
        assert layout["background"].upper() == "#000000"
        renderer._layout = layout
        assert renderer._bg_color().upper() == layout["background"].upper()

    deck = _pitch_like_deck()
    dc = DesignContext.from_profile(profile)
    pptx = renderer.render_deck(deck, dc)
    prs = Presentation(io.BytesIO(pptx))
    slide = list(prs.slides)[2]
    auditor.W, auditor.H = prs.slide_width, prs.slide_height
    text_shapes = [s for s in slide.shapes
                   if s.has_text_frame and s.text_frame.text.strip()]
    assert text_shapes
    for shape in text_shapes:
        bg = auditor._effective_bg(slide, shape)
        assert bg.upper() in ("#000000", "#032535", "#0077FF", "#DBECFF", "#FFFFFF")


def _pictures_deep(shapes):
    for shape in shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from _pictures_deep(shape.shapes)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                yield shape
        except Exception:  # noqa: BLE001
            continue


def _brightness(blob: bytes) -> float:
    from PIL import Image

    with Image.open(io.BytesIO(blob)) as img:
        rgba = img.convert("RGBA")
        if max(rgba.size) > 128:
            ratio = 128 / max(rgba.size)
            rgba = rgba.resize((max(1, round(rgba.width * ratio)),
                                max(1, round(rgba.height * ratio))), Image.NEAREST)
        pixels = [(r, g, b) for r, g, b, a in rgba.getdata() if a > 200]
    if not pixels:
        return 1.0
    return sum((r + g + b) / 3 for r, g, b in pixels) / len(pixels) / 255.0


# ---------------------------------------------------------------- Блок B: чеки
def _crafted_pptx(text_color: RGBColor, *, empty_box: bool = False) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "" if empty_box else "Невидимый текст"
    run.font.size = Pt(18)
    run.font.color.rgb = text_color
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _minimal_profile() -> dict:
    return {
        "slide_size": {"w_in": 10.0, "h_in": 7.5},
        "palette": [], "text_colors": [], "fonts": [],
        "type_scale": {}, "grid": {},
        "layouts": [{"id": "L0", "name": "Blank", "role": "content",
                     "kind": "blank", "role_reason": "тест", "score": 0.5,
                     "placeholders": [], "body": None, "columns": [],
                     "branding": [], "decor": [], "background": "#000000"}],
    }


def _minimal_deck() -> Deck:
    return Deck(title="Тёмный тест", slides=[
        Slide(slide_type=SlideType.TITLE, heading="Тёмный титул"),
        Slide(slide_type=SlideType.CONTENT, heading="Чёрный слайд",
              blocks=[Block(kind="bullets", items=["Тезис"])]),
        Slide(slide_type=SlideType.FINAL, heading="Финал презентации"),
    ])


def test_text_invisible_caught():
    """Чёрный текст на чёрном p:bg ловится как ошибка text_invisible."""
    pptx = _crafted_pptx(RGBColor(0, 0, 0))
    audit = Audit(_minimal_profile()).audit(_minimal_deck(), pptx)
    codes = {issue["code"] for issue in audit["issues"]}
    assert "text_invisible" in codes, audit["issues"]
    issue = next(i for i in audit["issues"] if i["code"] == "text_invisible")
    assert issue["severity"] == "error"
    assert issue["bbox"], "у проблемы есть координаты для подсветки"


def test_empty_text_frame_caught():
    """Пустая текстовая рамка рендера — предупреждение empty_text_frame."""
    pptx = _crafted_pptx(RGBColor(255, 255, 255), empty_box=True)
    audit = Audit(_minimal_profile()).audit(_minimal_deck(), pptx)
    codes = {issue["code"] for issue in audit["issues"]}
    assert "empty_text_frame" in codes, audit["issues"]
    issue = next(i for i in audit["issues"] if i["code"] == "empty_text_frame")
    assert issue["severity"] == "warning"
