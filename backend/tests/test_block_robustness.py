"""Устойчивость колоды к ошибкам модели в типах блоков.

Реальный случай на AITUNNEL: Qwen3.5 пометила блок `kind="factoids"`, а данные
положила в `items`. Блок считался пустым, слайд оставался с одним заголовком —
аудит справедливо ругал его как пустой, но причина была в схеме.
"""
from __future__ import annotations

import io

import pytest
from pptx import Presentation

from app.audit.checks import Audit
from app.layout.engine import LayoutEngine, DesignContext
from app.models.deck import Block, Deck, Slide, SlideType


def test_factoids_with_items_are_coerced():
    block = Block(kind="factoids", items=[
        "Время подготовки отчётов сократилось на 40%",
        "Автоматизировано 12 задач",
    ])
    assert block.items == []
    assert len(block.factoids) == 2
    assert block.factoids[0]["value"] == "40%"
    assert "Время" in block.factoids[0]["label"]
    assert block.factoids[1]["value"] == "12 задач"


def test_factoids_without_number_get_short_value():
    block = Block(kind="factoids", items=["Скорость принятия решений"])
    assert block.factoids[0]["value"] == "Скорость принятия"
    assert block.factoids[0]["label"] == "решений"


def test_table_without_table_is_rejected():
    """Явная ошибка лучше пустого слайда: планировщик повторит запрос."""
    with pytest.raises(Exception):
        Block(kind="table", items=["нет данных"])


def test_quote_text_in_text_field_is_moved():
    block = Block(kind="quote", text="Главное — скорость решений")
    assert block.quote_text == "Главное — скорость решений"
    assert block.text is None


@pytest.mark.parametrize("block,expected", [
    # kind назван неверно, но данные есть — виджет выбирается по данным
    (Block(kind="bullets", factoids=[{"value": "25%", "label": "рост"}]), "factoids"),
    (Block(kind="factoids", items=["Рост на 25%"]), "factoids"),  # схема привела данные
    (Block(kind="bullets", items=["Тезис"]), "bullets"),
    (Block(kind="steps", items=["Шаг"]), "steps"),
    (Block(kind="text", text="Абзац"), "text"),
    (Block(kind="quote", quote_text="Цитата"), "quote"),
    (Block(kind="image", image_ref="s1_img1.png"), "image"),
])
def test_widget_is_chosen_by_payload(block, expected):
    """kind — подсказка; виджет выбирается по фактическому содержимому."""
    assert LayoutEngine.widget_for(block) == expected


def test_numbered_widget_uses_real_autonumber(profile_of, render_variants,
                                               synthetic_template):
    """kind=numbered рисуется настоящей нумерацией PowerPoint, а не текстом «1.»."""
    deck = Deck(title="Проба нумерации", slides=[
        Slide(slide_type=SlideType.TITLE, heading="Титул презентации"),
        Slide(slide_type=SlideType.CONTENT, heading="Этапы реализации",
              blocks=[Block(kind="numbered", items=[
                  "Подключить 10 отделов", "Внедрить ML-предсказания",
                  "Обучить пользователей"])]),
        Slide(slide_type=SlideType.FINAL, heading="Финал презентации"),
    ])
    profile = profile_of(synthetic_template)
    pptx = render_variants(profile, deck, synthetic_template)["compact"]

    prs = Presentation(io.BytesIO(pptx))
    from pptx.oxml.ns import qn

    numbers = 0
    texts = []
    for shape in prs.slides[1].shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for paragraph in shape.text_frame.paragraphs:
            text = "".join(run.text for run in paragraph.runs)
            if not text.strip():
                continue
            texts.append(text)
            pPr = paragraph._p.find(qn("a:pPr"))
            if pPr is not None and pPr.find(qn("a:buAutoNum")) is not None:
                numbers += 1
    assert numbers == 3, f"ожидались три автонумерованных пункта: {texts}"
    assert "1." not in " ".join(texts), "номер не должен быть частью текста"
    assert LayoutEngine.widget_for(Block(kind="numbered", items=["раз"])) == "numbered"


def test_empty_block_detection_is_payload_based():
    assert LayoutEngine._empty(Block(kind="bullets", items=["Тезис"])) is False
    assert LayoutEngine._empty(Block(kind="factoids", items=["Рост на 25%"])) is False
    assert LayoutEngine._empty(Block(kind="text", text="Абзац")) is False
    assert LayoutEngine._empty(Block(kind="factoids")) is True


def test_mislabeled_block_still_renders(profile_of, render_variants, synthetic_template):
    """Слайд с «неправильным» типом блока не должен оставаться пустым."""
    deck = Deck(title="Проба типов", slides=[
        Slide(slide_type=SlideType.TITLE, heading="Титул презентации"),
        Slide(slide_type=SlideType.CONTENT, heading="Метрики квартала",
              blocks=[Block(kind="factoids", items=[
                  "Время отчётов сократилось на 40%",
                  "Охват вырос до 5 подразделений",
                  "2000 сотрудников используют платформу"])]),
        Slide(slide_type=SlideType.FINAL, heading="Финал презентации"),
    ])
    profile = profile_of(synthetic_template)
    pptx = render_variants(profile, deck, synthetic_template)["compact"]

    prs = Presentation(io.BytesIO(pptx))
    content_slide = prs.slides[1]
    texts = [shape.text_frame.text for shape in content_slide.shapes
             if shape.has_text_frame and shape.text_frame.text.strip()]
    assert any("40%" in text for text in texts), texts

    audit = Audit(profile).audit(deck, pptx)
    assert "slide_too_sparse" not in {issue["code"] for issue in audit["issues"]}
