"""Сквозные тесты пайплайна на синтетическом шаблоне (работают без файлов VK)."""
import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from app.content.importer import import_content_pack
from app.pipeline import build_profile, full_generate

BRIEF = ("Хакатон VK Tech: нужна презентация проекта «Цифровой дизайнер презентаций». "
         "Решение генерирует колоду по брифу на произвольном шаблоне, верстает 3 варианта "
         "нативными объектами, прогоняет детерминированный и VLM-аудит и возвращает "
         "pptx/pdf/html. Конкурсное требование — открытые модели, без платных API.")


def test_full_generate_three_variants_without_llm(synthetic_template):
    res = full_generate(BRIEF, "", "project", synthetic_template, "synthetic.pptx")
    assert res["planner"]["used_llm"] is False
    assert len(res["variants"]) == 3
    for variant in res["variants"]:
        assert variant["name"] in ("compact", "cards", "split")
        assert variant["pptx"].startswith(b"PK")
        assert variant["audit"]["passed"], variant["audit"]["issues"][:2]
    assert "<!DOCTYPE html>" in res["html"]
    assert res["deck"]["slides"]
    # VLM недоступен в тестовом окружении — стадия обязана деградировать, а не падать
    assert res["vlm"]["available"] is False


def test_profile_is_purpose_agnostic(synthetic_template):
    profile = build_profile(synthetic_template, "x.pptx")
    assert profile["template_id"]
    assert profile["slide_size"]["w_in"] > 0
    assert profile["layouts"]
    assert profile["palette"]
    assert profile["headline_font"]


def test_generate_with_content_pack_embeds_images(synthetic_template, tiny_png):
    """Контент-пакет: текст идёт в колоду, изображения — в слайды нативными объектами."""
    prs = Presentation(io.BytesIO(synthetic_template))
    prs.slides[0].shapes.add_picture(io.BytesIO(tiny_png), Inches(1), Inches(1),
                                     width=Inches(2), height=Inches(1))
    buf = io.BytesIO()
    prs.save(buf)
    corpus = import_content_pack(buf.getvalue(), "pack.pptx")
    assert corpus.images

    res = full_generate(BRIEF, "", "project", synthetic_template, "synthetic.pptx",
                        corpus=corpus)
    assert res["corpus"]["stats"]["images"] == len(corpus.images)
    assert res["planner"]["used_llm"] is False
    image_blocks = [b for slide in res["deck"]["slides"] for b in slide["blocks"]
                    if b["kind"] == "image"]
    assert image_blocks, "в колоде должен появиться блок-иллюстрация из корпуса"
    assert image_blocks[0]["image_ref"] in corpus.images

    found_pictures = False
    for variant in res["variants"]:
        assert variant["audit"]["errors"] == 0, variant["audit"]["issues"][:3]
        prs_out = Presentation(io.BytesIO(variant["pptx"]))
        for slide in prs_out.slides:
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    found_pictures = True
    assert found_pictures, "изображение из контент-пакета должно быть встроено нативно"


def test_generate_on_unfamiliar_template(unfamiliar_template):
    """Мутированный шаблон (имена макетов Slide N, чужие токены) обязан работать."""
    res = full_generate(BRIEF, "", "project", unfamiliar_template, "unfamiliar.pptx")
    profile = res["profile"]
    assert profile["layouts"]
    roles = profile["layout_groups"]
    assert "content" in roles, f"роли макетов не определены: {roles}"
    for variant in res["variants"]:
        assert variant["audit"]["errors"] == 0, variant["audit"]["issues"][:3]
