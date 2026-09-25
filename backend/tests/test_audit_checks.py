"""Негативные тесты детерминированных проверок: каждая ловит свой дефект.

Приём один: берём корректно сгенерированную колоду, вносим ровно один дефект
(или аудитуем с чужим профилем) и проверяем, что нужный код проблемы появился.
"""
from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from app.audit.checks import Audit
from app.planner.fallback import FallbackPlanner

BRIEF = ("Платформа аналитики сократила время подготовки отчётов на 40%, "
         "автоматизировала 12 задач, охватила 5 подразделений. 2000 сотрудников "
         "используют её еженедельно. План: подключить 10 отделов и внедрить "
         "ML-предсказания выручки.")


@pytest.fixture(scope="module")
def deck():
    return FallbackPlanner().plan(BRIEF, "", "project")


# в колоде офлайн-планировщика: 0 — титул, 1 — оглавление, 2 — раздел,
# 3.. — контентные слайды. Проверки плотности применяются только к контенту.
CONTENT_SLIDE = 3


@pytest.fixture(scope="module")
def rendered(profile_of, render_variants, synthetic_template, deck):
    """Колода «compact» на синтетическом шаблоне: базис без ошибок до мутаций.

    Допускаются замечания `text_overflow`: офлайн-планировщик кладёт на слайд
    три блока (тезисы + фактоиды + диаграмма), и в компактной вёрстке текстовый
    блок получает меньше высоты, чем нужно шкале шаблона. Это честное замечание
    аудита (визуально текст уходит в зазор между блоками), а не ошибка. Ошибки
    по-прежнему запрещены: тесты мутируют именно чистый базис.
    """
    profile = profile_of(synthetic_template)
    pptx = render_variants(profile, deck, synthetic_template)["compact"]
    baseline = Audit(profile).audit(deck, pptx)
    assert baseline["errors"] == 0, "базис должен быть без ошибок"
    assert {i["code"] for i in baseline["issues"]} <= {"text_overflow"}, \
        f"неожиданные замечания базиса: {[i['code'] for i in baseline['issues']]}"
    return pptx


def _codes_after(profile: dict, deck, pptx: bytes, mutate) -> set[str]:
    prs = Presentation(io.BytesIO(pptx))
    mutate(prs)
    buf = io.BytesIO()
    prs.save(buf)
    result = Audit(profile).audit(deck, buf.getvalue())
    return {issue["code"] for issue in result["issues"]}


def _textbox(prs, text: str, *, x=1.0, y=3.0, w=4.0, h=0.6, size=18, font=None,
             name=None):
    shape = prs.slides[1].shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    if font:
        run.font.name = font
    if name:
        shape.name = name
    return shape


# ------------------------------------------------------------------ шаблон
def test_font_size_not_in_scale(profile_of, synthetic_template, deck, rendered):
    profile = profile_of(synthetic_template)
    codes = _codes_after(profile, deck, rendered,
                         lambda prs: _textbox(prs, "Мелкий текст вне шкалы", size=7))
    assert "font_size_not_in_scale" in codes


def test_too_many_typefaces(profile_of, synthetic_template, deck, rendered):
    profile = profile_of(synthetic_template)
    codes = _codes_after(profile, deck, rendered,
                         lambda prs: _textbox(prs, "Третий шрифт в колоде",
                                              size=18, font="Comic Sans MS"))
    assert "too_many_typefaces" in codes


def test_layout_not_from_template(profile_of, synthetic_template, deck, rendered):
    """Аудит с чужим профилем видит, что макеты слайда не из этого шаблона."""
    from tools.make_fixtures import FIXTURES, build_template

    other_profile = profile_of(build_template(**FIXTURES["synthetic_4x3"]))
    other_profile = dict(other_profile)
    other_profile["layouts"] = [
        dict(layout, name=f"Чужой макет {i}")
        for i, layout in enumerate(other_profile["layouts"])]
    result = Audit(other_profile).audit(deck, rendered)
    assert "layout_not_from_template" in {i["code"] for i in result["issues"]}


def test_branding_shifted(profile_of, synthetic_template, deck, rendered):
    profile = profile_of(synthetic_template)
    # имена фирменных элементов зависят от макета, поэтому берём их из профиля
    # именно того макета, на котором стоит правимый слайд
    layout_name = Presentation(io.BytesIO(rendered)).slides[1].slide_layout.name
    layout = next(item for item in profile["layouts"] if item["name"] == layout_name)
    brand = layout["branding"][0]

    def mutate(prs):
        _textbox(prs, "Колонтитул", x=brand["x"] + 0.6, y=brand["y"],
                 w=brand["w"], h=brand["h"], size=12, name=brand["name"])

    assert "branding_shifted" in _codes_after(profile, deck, rendered, mutate)


def test_misaligned_to_grid(profile_of, synthetic_template, deck, rendered):
    """Два блока одной колонки не выровнены на 0.05″."""
    profile = profile_of(synthetic_template)

    def mutate(prs):
        first = _textbox(prs, "Первый блок колонки", x=1.2, y=2.0, w=4.0, h=0.5)
        _textbox(prs, "Второй блок колонки", x=1.25, y=2.6, w=4.0, h=0.5)
        assert first.width == Inches(4.0)

    assert "misaligned_to_grid" in _codes_after(profile, deck, rendered, mutate)


def test_text_overflow(profile_of, synthetic_template, deck, rendered):
    """Текст больше своей рамки: 30 pt в коробке 0.18″ (случай factoid-подписей).

    Регрессия: проверка сравнивала дюймы с EMU и не срабатывала никогда, из-за
    чего наложение подписей factoid'ов проходило аудит как «чисто».
    """
    profile = profile_of(synthetic_template)
    codes = _codes_after(profile, deck, rendered,
                         lambda prs: _textbox(prs, "150 000 человек", x=1.0, y=3.0,
                                              w=4.0, h=0.18, size=30))
    assert "text_overflow" in codes


def test_content_in_margins(profile_of, synthetic_template, deck, rendered):
    profile = profile_of(synthetic_template)
    codes = _codes_after(profile, deck, rendered,
                         lambda prs: _textbox(prs, "Текст у самого края", x=0.05, y=3.0))
    assert "content_in_margins" in codes


# ------------------------------------------------------------------ вёрстка
def test_image_stretched(profile_of, synthetic_template, deck, rendered, tiny_png):
    """Картинка 2:1 растянута в квадрат."""
    profile = profile_of(synthetic_template)

    def mutate(prs):
        prs.slides[1].shapes.add_picture(io.BytesIO(tiny_png), Inches(1), Inches(3),
                                         width=Inches(3), height=Inches(3))

    assert "image_stretched" in _codes_after(profile, deck, rendered, mutate)


def test_slide_too_sparse(profile_of, synthetic_template, deck, rendered):
    """Со слайда убрано всё, кроме заголовка."""
    profile = profile_of(synthetic_template)

    def mutate(prs):
        slide = prs.slides[CONTENT_SLIDE]
        for shape in list(slide.shapes):
            if not (shape.is_placeholder and shape.has_text_frame
                    and shape.text_frame.text.strip()):
                shape._element.getparent().remove(shape._element)

    assert "slide_too_sparse" in _codes_after(profile, deck, rendered, mutate)


def test_slide_too_dense(profile_of, synthetic_template, deck, rendered):
    """Фоновая фигура на 80% площади слайда считается переполнением."""
    profile = profile_of(synthetic_template)

    def mutate(prs):
        from pptx.enum.shapes import MSO_SHAPE
        prs.slides[CONTENT_SLIDE].shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.4), Inches(12.5), Inches(6.7))

    assert "slide_too_dense" in _codes_after(profile, deck, rendered, mutate)


# ------------------------------------------------------------------ данные
def test_chart_unlabeled(profile_of, synthetic_template, deck, rendered):
    profile = profile_of(synthetic_template)

    def mutate(prs):
        data = CategoryChartData()
        data.categories = ["Q1", "Q2"]
        data.add_series("Без подписей", (1.0, 2.0))
        prs.slides[1].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1),
                                       Inches(3), Inches(6), Inches(3), data)

    assert "chart_unlabeled" in _codes_after(profile, deck, rendered, mutate)


def test_duplicate_slide(profile_of, synthetic_template, render_variants, deck):
    """Две одинаковые по содержанию слайда — дубли."""
    from app.models.deck import Block, Deck, Slide, SlideType

    duplicated = Deck(title="Дубли", slides=[
        Slide(slide_type=SlideType.TITLE, heading="Титул презентации"),
        Slide(slide_type=SlideType.CONTENT, heading="Одинаковый слайд",
              blocks=[Block(kind="bullets", items=["Тезис один", "Тезис два"])]),
        Slide(slide_type=SlideType.CONTENT, heading="Одинаковый слайд",
              blocks=[Block(kind="bullets", items=["Тезис один", "Тезис два"])]),
        Slide(slide_type=SlideType.FINAL, heading="Финал презентации"),
    ])
    profile = profile_of(synthetic_template)
    pptx = render_variants(profile, duplicated, synthetic_template)["compact"]
    result = Audit(profile).audit(duplicated, pptx)
    codes = {i["code"] for i in result["issues"]}
    assert "duplicate_slide" in codes
    assert "duplicate_heading" in codes


def test_fill_color_not_allowed(profile_of, synthetic_template, deck, rendered):
    """Случайная заливка ловится, а смесь палитры (светлая подложка) — нет."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    profile = profile_of(synthetic_template)
    audit = Audit(profile)
    palette = [p["hex"] for p in profile["palette"]]

    def mutate(prs):
        alien = prs.slides[1].shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(1), Inches(3), Inches(2), Inches(1))
        alien.fill.solid()
        alien.fill.fore_color.rgb = RGBColor(0x12, 0xFF, 0x34)

    codes = _codes_after(profile, deck, rendered, mutate)
    assert "color_not_allowed" in codes

    # смесь двух цветов палитры разрешена (ADR-008)
    assert audit.is_derived_color(palette[1])
    assert not audit.is_derived_color("#12FF34")


def test_issues_have_ids_and_bbox(profile_of, synthetic_template, deck, rendered):
    """Каждая проблема пригодна для UI: есть id, severity, bbox и признак детерминизма."""
    profile = profile_of(synthetic_template)
    result = Audit(profile).audit(deck, rendered)
    for issue in result["issues"]:
        assert issue["id"], "у проблемы должен быть стабильный id"
        assert issue["severity"] in ("error", "warning")
        assert isinstance(issue["bbox"], list)
        assert issue["deterministic"] is True
