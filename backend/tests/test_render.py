import io

from pptx import Presentation

from app.layout.engine import LayoutEngine
from app.planner.fallback import FallbackPlanner
from app.render.pptx_renderer import Renderer

BRIEF = ("Компания запускает мобильное приложение доставки еды. Охват — 3 города, "
         "150 000 пользователей, средний чек 870 руб., доставка быстрее на 18%. "
         "Цель — 10 городов к концу года и рост LTV на 25%.")
META = {
    "layouts": [{"name": "Титул"}, {"name": "Слайд с заголовком"}],
    "masters": [{"name": "Офисное"}],
}


def test_render_roundtrip(template_any):
    deck = FallbackPlanner().plan(BRIEF, "", "project")
    profile = build_profile(template_any)
    renderer = Renderer(profile)
    for variant in ("compact", "cards", "split"):
        engine = LayoutEngine(profile, deck, variant)
        slides = engine.build()
        pptx = renderer.render(deck, slides)
        assert isinstance(pptx, bytes) and pptx.startswith(b"PK")
        prs = Presentation(io.BytesIO(pptx))
        assert len(list(prs.slides)) == len(deck.slides)
        # каждый слайд содержит хоть одну фигуру с текстом
        for slide in prs.slides:
            texts = [sh for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
            assert texts, f"пустой слайд в варианте {variant}"


def build_profile(template_bytes: bytes) -> dict:
    from app.template.parser import TemplateParser

    return TemplateParser(template_bytes).parse().to_dict()