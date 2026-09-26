"""HTML+CSS → PPTX: только нативные фигуры, без растрового слайда (ADR-036)."""
from __future__ import annotations

import io
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.audit.checks import Audit
from app.models.deck import Block, Chart, ChartType, Deck, Slide, SlideType, Table
from app.render.html_renderer import render_html
from app.render.html_to_pptx import html_to_pptx

BRIEF = ("Платформа аналитики сократила время подготовки отчётов на 40%, "
         "автоматизировала 12 задач, охватила 5 подразделений.")

FACTOIDS = [
    {"value": "180", "label": "автоматических тестов"},
    {"value": "24/24", "label": "чистых комбинаций"},
    {"value": "5 мин", "label": "бюджет задания"},
]
STEPS = ["Пилот в подразделении.", "Два корпоративных шаблона.", "Масштабирование."]


def _deck() -> Deck:
    return Deck(title="Конвертер HTML в PPTX", slides=[
        Slide(slide_type=SlideType.TITLE, heading="Титул презентации",
              subheading="Итоги квартала"),
        Slide(slide_type=SlideType.CONTENT, heading="Тезисы проекта",
              blocks=[Block(kind="bullets", items=["Первый тезис.", "Второй тезис."])]),
        Slide(slide_type=SlideType.CONTENT, heading="Метрики проекта",
              blocks=[Block(kind="factoids", factoids=FACTOIDS)]),
        Slide(slide_type=SlideType.CONTENT, heading="Сравнение подходов",
              blocks=[Block(kind="table", table=Table(
                  header=["Критерий", "Ручная", "Сервис"],
                  rows=[["Время", "8 ч", "1 мин"], ["Ошибки", "15%", "0%"]]))]),
        Slide(slide_type=SlideType.CONTENT, heading="Динамика внедрения",
              blocks=[Block(kind="chart", chart=Chart(
                  type=ChartType.COLUMN, categories=["Q1", "Q2", "Q3"], unit="%",
                  series=[{"name": "Охват", "values": [20.0, 45.0, 80.0]}]))]),
        Slide(slide_type=SlideType.CONTENT, heading="План внедрения",
              blocks=[Block(kind="steps", items=STEPS)]),
        Slide(slide_type=SlideType.FINAL, heading="Спасибо за внимание"),
    ])


def _convert(profile: dict, template: bytes, variant: str = "compact") -> tuple[str, bytes]:
    html = render_html(_deck(), profile, variant=variant, template_bytes=template)
    return html, html_to_pptx(html, profile, template_bytes=template)


def test_html_to_pptx_native_objects(profile_of, synthetic_template):
    profile = profile_of(synthetic_template)
    _, pptx = _convert(profile, synthetic_template)
    prs = Presentation(io.BytesIO(pptx))
    counts = {"text_frames": 0, "tables": 0, "charts": 0, "shapes": 0, "pictures": 0}
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                counts["pictures"] += 1
            elif getattr(shape, "has_table", False) and shape.has_table:
                counts["tables"] += 1
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                counts["charts"] += 1
            elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                counts["shapes"] += 1
                if shape.has_text_frame and shape.text_frame.text.strip():
                    counts["text_frames"] += 1
            elif shape.has_text_frame and shape.text_frame.text.strip():
                counts["text_frames"] += 1
    assert counts["text_frames"] >= 10, counts
    assert counts["tables"] >= 1, counts
    assert counts["charts"] >= 1, counts
    assert counts["shapes"] >= 3, "шевроны шагов должны быть нативными фигурами"
    assert counts["pictures"] == 0, "картинок в этой колоде нет"


def test_html_to_pptx_no_full_slide_picture(profile_of, template_workspace):
    profile = profile_of(template_workspace)
    _, pptx = _convert(profile, template_workspace)
    prs = Presentation(io.BytesIO(pptx))
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            assert not (shape.width >= prs.slide_width * 0.92
                        and shape.height >= prs.slide_height * 0.92), \
                "слайд не должен быть растровой картинкой"


def test_html_to_pptx_preserves_colors(profile_of, synthetic_template):
    profile = profile_of(synthetic_template)
    html, pptx = _convert(profile, synthetic_template)
    root = re.search(r":root \{(.*?)\}", html, re.S).group(1)
    tokens = dict(piece.split(":", 1) for piece in root.split(";") if ":" in piece)
    foreground = tokens["--color-text"].strip().lstrip("#").upper()
    # заливка шапки берётся из самой таблицы: она может отличаться от корневого
    # акцента (виджет использует акцент DesignContext) — важно, что CSS→PPTX
    # переносит цвет без потерь
    table_style = re.search(r"<table[^>]*style=\"([^\"]*)\"", html).group(1)
    table_tokens = dict(piece.split(":", 1) for piece in table_style.split(";") if ":" in piece)
    header_bg = table_tokens["--header-bg"].strip().lstrip("#").upper()

    prs = Presentation(io.BytesIO(pptx))
    colors = []
    header_fills = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        try:
                            if run.font.color and run.font.color.type is not None:
                                colors.append(str(run.font.color.rgb).upper())
                        except Exception:  # noqa: BLE001
                            continue
            if getattr(shape, "has_table", False) and shape.has_table:
                cell = shape.table.cell(0, 0)
                try:
                    header_fills.append(str(cell.fill.fore_color.rgb).upper())
                except Exception:  # noqa: BLE001
                    continue
    assert foreground in colors, f"цвет текста {foreground} не найден среди {set(colors)}"
    assert header_bg in header_fills, \
        f"заливка шапки {header_bg} не найдена среди {header_fills}"


def test_html_to_pptx_roundtrip(profile_of, template_workspace):
    profile = profile_of(template_workspace)
    deck = _deck()
    html, pptx = _convert(profile, template_workspace)
    prs = Presentation(io.BytesIO(pptx))
    assert len(list(prs.slides)) == len(deck.slides)

    all_text = []
    chart_categories = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                all_text.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    all_text.extend(cell.text.strip() for cell in row.cells)
            if getattr(shape, "has_chart", False) and shape.has_chart:
                chart_categories.extend(list(shape.chart.plots[0].categories))
    joined = "\n".join(all_text)
    for fact in FACTOIDS:
        assert fact["value"] in joined, fact
        assert fact["label"] in joined, fact
    for step in STEPS:
        assert step in joined, step
    for header in ("Критерий", "Ручная", "Сервис"):
        assert header in joined, header
    assert chart_categories == ["Q1", "Q2", "Q3"], chart_categories


def test_audit_runs_on_both_renderers(profile_of, template_workspace):
    """Один и тот же детерминированный аудит применяется к обоим PPTX."""
    from app.layout.engine import DesignContext
    from app.render.pptx_renderer import Renderer

    profile = profile_of(template_workspace)
    deck = _deck()
    dc = DesignContext.from_profile(profile)
    native = Renderer(profile, variant="compact",
                      template_bytes=template_workspace).render_deck(deck, dc)
    html, converted = _convert(profile, template_workspace, variant="compact")
    for label, pptx in (("native", native), ("html", converted)):
        audit = Audit(profile).audit(deck, pptx)
        assert audit["errors"] == 0, f"{label}: {audit['issues'][:3]}"
    assert "data-widget" in html
