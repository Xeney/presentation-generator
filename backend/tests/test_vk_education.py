"""Регрессия VK Education: обрезанный и невидимый текст на панелях декора.

Реальный сбой: на шаблоне VK Education 7 из 13 слайдов выходили «обрывками» —
вёрстка брала макет «Заголовок в 2 строки + объект» (левая половина под текст,
правая — тёмная панель) и рисовала текст на всю ширину слайда. Часть строк
попадала на тёмную панель и перекрашивалась в белый, а на белом фоне пропадала:
аудит при этом молчал (0 ошибок), поэтому проверка усилена.
"""
from __future__ import annotations

import io

from pptx import Presentation

from app.audit.checks import Audit
from app.layout.engine import DesignContext, avoid_decor_panels
from app.layout.geometry import Rect
from app.models.deck import Block, Deck, Slide, SlideType
from app.render.pptx_renderer import Renderer

VARIANTS = ("compact", "cards", "split")
BROKEN_CODES = {"text_overflow", "text_on_decor_edge", "overlap",
                "out_of_bounds", "content_in_margins"}


def _pitch_deck() -> Deck:
    """Колода с теми же композициями, что ломались в живом прогоне."""
    bullets = [("Ключевые принципы подхода", [
        "Шаблон — это набор данных, а не статичный файл.",
        "LLM отвечает за смысл контента, а не за вёрстку.",
        "Аудит является частью пайплайна, а не внешней проверкой.",
        "Использование открытых весов для локальной работы.",
        "Гарантия воспроизведения дизайн-системы любого файла.",
    ])]
    slides = [
        Slide(slide_type=SlideType.TITLE, heading="Проект: AI-генерация презентаций",
              subheading="Питч проекта"),
        Slide(slide_type=SlideType.CONTENT, heading="Архитектура пайплайна",
              blocks=[Block(kind="steps", items=[
                  "Парсинг шаблона в JSON-профиль.",
                  "Формирование колоды локальной моделью.",
                  "Планирование контента.",
                  "Вёрстка трёх вариантов.",
                  "Рендер нативными объектами PowerPoint.",
                  "Аудит и авто-фиксы.",
              ])]),
        Slide(slide_type=SlideType.CONTENT, heading="Метрики и результаты тестирования",
              blocks=[Block(kind="factoids", factoids=[
                  {"value": "180", "label": "выполненных тестов"},
                  {"value": "8", "label": "шаблонов в матрице"},
                  {"value": "3", "label": "варианта вёрстки"},
                  {"value": "24 из 24", "label": "чистых комбинаций"},
              ])]),
    ]
    slides += [
        Slide(slide_type=SlideType.CONTENT, heading=heading,
              blocks=[Block(kind="bullets", items=items)])
        for heading, items in bullets
    ]
    slides.append(Slide(slide_type=SlideType.FINAL, heading="Следующие шаги",
                        blocks=[Block(kind="bullets", items=["Пилот в подразделении."])]))
    return Deck(title="Питч проекта", slides=slides)


def _frames(pptx: bytes) -> list[tuple]:
    prs = Presentation(io.BytesIO(pptx))
    W, H = prs.slide_width, prs.slide_height
    out = []
    for si, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.left is None or shape.width is None:
                continue
            out.append((si, shape, shape.left, shape.top or 0, shape.width,
                        shape.height or 0, W, H))
    return out


def test_vk_education_no_overflow(profile_of, template_edu):
    """Все текстовые фреймы внутри границ, текст не обрезан и не исчезает."""
    profile = profile_of(template_edu)
    dc = DesignContext.from_profile(profile)
    deck = _pitch_deck()
    renderer = Renderer(profile, variant="compact", template_bytes=template_edu)

    # область контента не заходит на цветные панели макета (причина обрывков)
    for slide in deck.slides:
        layout = renderer._pick_layout(slide.slide_type, slide)
        canvas = renderer._canvas(layout, slide)
        for item in layout.get("decor") or []:
            left = max(canvas.x, float(item.get("x", 0)))
            top = max(canvas.y, float(item.get("y", 0)))
            right = min(canvas.right, float(item.get("x", 0)) + float(item.get("w", 0)))
            bottom = min(canvas.bottom, float(item.get("y", 0)) + float(item.get("h", 0)))
            if right <= left or bottom <= top:
                continue
            overlap = (right - left) * (bottom - top)
            assert overlap < 0.25 * canvas.w * canvas.h, (
                f"{slide.heading}: текстовая область заходит на панель "
                f"«{item.get('name')}» ({item.get('fill')})")

    for variant in VARIANTS:
        renderer = Renderer(profile, variant=variant, template_bytes=template_edu)
        pptx = renderer.render_deck(deck, dc)
        audit = Audit(profile).audit(deck, pptx)
        codes = {issue["code"] for issue in audit["issues"]}
        assert audit["errors"] == 0, codes
        assert not (codes & BROKEN_CODES), codes
        # фреймы физически внутри слайда
        for si, shape, x, y, w, h, W, H in _frames(pptx):
            assert x >= -10000 and y >= -10000, \
                f"{variant} слайд {si + 1}: «{shape.name}» вышел за левый/верхний край"
            assert x + w <= W + 10000 and y + h <= H + 10000, \
                f"{variant} слайд {si + 1}: «{shape.name}» вышел за правый/нижний край"


def test_choose_canvas_avoids_decor_panel(profile_of, template_edu):
    """Макет с тёмной половиной: канвас текста не пересекает панель."""
    profile = profile_of(template_edu)
    half = [l for l in profile["layouts"]
            if l.get("kind") == "bullets" and l.get("decor")]
    assert half, "ожидается макет с декор-панелью"
    layout = half[0]
    rect = avoid_decor_panels(Rect(0.73, 1.94, 11.88, 4.88), layout,
                              profile["slide_size"]["w_in"], profile["slide_size"]["h_in"])
    for item in layout["decor"]:
        x, y = float(item["x"]), float(item["y"])
        w, h = float(item["w"]), float(item["h"])
        left, top = max(rect.x, x), max(rect.y, y)
        right = min(rect.right, x + w)
        bottom = min(rect.bottom, y + h)
        inter = max(0.0, right - left) * max(0.0, bottom - top)
        assert inter < 0.25 * rect.w * rect.h
    assert rect.w >= 4.0, "свободной области должно хватать под текст"


def test_title_blocks_do_not_overlap(profile_of, template_edu):
    """Два блока на витринном слайде раскладываются, а не печатаются друг на друге."""
    profile = profile_of(template_edu)
    deck = Deck(title="Титул проекта", slides=[
        Slide(slide_type=SlideType.TITLE, heading="Проект: генератор презентаций",
              blocks=[Block(kind="text", text="Проблема: ручная переверстка."),
                      Block(kind="text", text="Решение: шаблон как данные.")]),
        Slide(slide_type=SlideType.CONTENT, heading="Содержательный слайд",
              blocks=[Block(kind="bullets", items=["Тезис"])]),
        Slide(slide_type=SlideType.FINAL, heading="Финал презентации"),
    ])
    dc = DesignContext.from_profile(profile)
    for variant in VARIANTS:
        renderer = Renderer(profile, variant=variant, template_bytes=template_edu)
        pptx = renderer.render_deck(deck, dc)
        audit = Audit(profile).audit(deck, pptx)
        overlaps = [i for i in audit["issues"] if i["code"] == "overlap"]
        assert not overlaps, overlaps
