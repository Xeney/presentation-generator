"""Тесты детерминированного аудита: позитивный кейс и негативные проверки.

Негативные кейсы строятся мутацией уже сгенерированного PPTX: так проверяется
именно детектор, а не схема Pydantic (которая не даст создать невалидный Deck).
"""
from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.audit.checks import Audit
from app.planner.fallback import FallbackPlanner

BRIEF = ("Платформа аналитики VK Tech сократила время подготовки отчётов на 40%, "
         "автоматизировала 12 задач и охватила 5 подразделений. 2000 сотрудников "
         "используют её еженедельно. План — 10 отделов к концу года и ML-предсказания "
         "выручки.")


@pytest.fixture(scope="module")
def deck():
    return FallbackPlanner().plan(BRIEF, "", "project")


def _reopen(pptx: bytes) -> Presentation:
    return Presentation(io.BytesIO(pptx))


def _save(prs: Presentation) -> bytes:
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_audit_positive_case(profile_of, render_variants, synthetic_template, deck):
    """Сгенерированные колоды чисты во всех вариантах: ошибок нет.

    Регрессионный тест на «одинаково валидны»: варианты различаются вёрсткой, но
    ни один не должен давать ошибок. Единственное допустимое замечание —
    `text_overflow`: офлайн-планировщик ставит на слайд три блока (тезисы +
    фактоиды + диаграмма), и в компактной вёрстке текстовый блок получает меньше
    высоты, чем требует шкала шаблона. Текст уходит в зазор между блоками, но
    аудит обязан об этом сказать (раньше проверка не срабатывала никогда —
    сравнивала дюймы с EMU). Убрать замечание можно только вёрсткой: это задача
    слоя layout (см. docs/DECISIONS.md, этап D).
    """
    profile = profile_of(synthetic_template)
    variants = render_variants(profile, deck, synthetic_template)
    for variant, pptx in variants.items():
        result = Audit(profile).audit(deck, pptx)
        assert result["passed"], f"{variant}: {result['issues'][:4]}"
        assert result["errors"] == 0, f"{variant}: {result['issues'][:4]}"
        extra = {i["code"] for i in result["issues"]} - {"text_overflow"}
        assert not extra, f"{variant}: {[i for i in result['issues'] if i['code'] in extra]}"


def test_audit_deterministic_and_ids_unique(profile_of, render_variants,
                                            synthetic_template, deck):
    profile = profile_of(synthetic_template)
    pptx = render_variants(profile, deck, synthetic_template)["compact"]
    first = Audit(profile).audit(deck, pptx)
    second = Audit(profile).audit(deck, pptx)
    assert first["issues"] == second["issues"], "аудит должен быть детерминированным"
    ids = [i["id"] for i in first["issues"]]
    assert len(ids) == len(set(ids)), "идентификаторы проблем должны быть уникальны"


def test_audit_catches_out_of_bounds(profile_of, render_variants,
                                     synthetic_template, deck):
    profile = profile_of(synthetic_template)
    prs = _reopen(render_variants(profile, deck, synthetic_template)["compact"])
    box = prs.slides[1].shapes.add_textbox(Inches(-4.5), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "уехавший блок"
    result = Audit(profile).audit(deck, _save(prs))
    codes = {i["code"] for i in result["issues"]}
    assert "out_of_bounds" in codes
    issue = next(i for i in result["issues"] if i["code"] == "out_of_bounds")
    assert issue["severity"] == "error"
    assert len(issue["bbox"]) == 4
    assert issue["deterministic"] is True


def test_audit_catches_table_too_big(profile_of, render_variants,
                                     synthetic_template, deck):
    """Регрессия: проверка размера таблицы была недостижима и не срабатывала."""
    profile = profile_of(synthetic_template)
    prs = _reopen(render_variants(profile, deck, synthetic_template)["compact"])
    slide = prs.slides[1]
    table = slide.shapes.add_table(8, 6, Inches(0.5), Inches(1.5),
                                   Inches(6), Inches(4)).table
    for r in range(8):
        for c in range(6):
            table.cell(r, c).text = f"{r}-{c}"
    result = Audit(profile).audit(deck, _save(prs))
    codes = {i["code"] for i in result["issues"]}
    assert "table_too_big" in codes, [i["code"] for i in result["issues"]]


def test_audit_catches_too_many_series(profile_of, render_variants,
                                       synthetic_template, deck):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    profile = profile_of(synthetic_template)
    prs = _reopen(render_variants(profile, deck, synthetic_template)["compact"])
    data = CategoryChartData()
    data.categories = ["A", "B", "C"]
    for n in range(6):
        data.add_series(f"S{n}", (1.0 + n, 2.0 + n, 3.0 + n))
    prs.slides[1].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5),
                                   Inches(6), Inches(3.5), data)
    result = Audit(profile).audit(deck, _save(prs))
    assert "too_many_series" in {i["code"] for i in result["issues"]}


def test_audit_catches_placeholder_text(profile_of, render_variants,
                                        synthetic_template, deck):
    profile = profile_of(synthetic_template)
    prs = _reopen(render_variants(profile, deck, synthetic_template)["compact"])
    box = prs.slides[1].shapes.add_textbox(Inches(0.5), Inches(6), Inches(4), Inches(0.6))
    box.text_frame.text = "TODO: вставьте текст"
    result = Audit(profile).audit(deck, _save(prs))
    assert "placeholder_text" in {i["code"] for i in result["issues"]}


def test_audit_catches_raster_slide(profile_of, render_variants, synthetic_template,
                                    deck, tiny_png):
    profile = profile_of(synthetic_template)
    prs = _reopen(render_variants(profile, deck, synthetic_template)["compact"])
    prs.slides[1].shapes.add_picture(io.BytesIO(tiny_png), 0, 0,
                                     width=prs.slide_width, height=prs.slide_height)
    result = Audit(profile).audit(deck, _save(prs))
    assert "raster_slide" in {i["code"] for i in result["issues"]}


def test_audit_rejects_broken_file(profile_of, synthetic_template, deck):
    profile = profile_of(synthetic_template)
    with pytest.raises(Exception):
        Audit(profile).audit(deck, b"not a pptx at all")
