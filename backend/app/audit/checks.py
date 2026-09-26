"""Детерминированный аудит сгенерированной PPTX по гайдлайнам шаблона.

Работает на python-pptx напрямую: границы, наложения, обрезка текста,
разрешённые токены (шрифты/цвета), контраст WCAG (>=4.5:1), плотность контента,
пустые слайды и заглушки, дубликаты, запрет «слайда-картинки».
Возвращает список проблем с bbox (в долях слайда) для красных рамок.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field, asdict

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from ..content.corpus import looks_like_placeholder
from ..layout.engine import choose_canvas, slide_needs_uniform_background
from ..models.deck import Deck, SlideType
from ..render.images import hex_to_rgb
from ..render.pptx_renderer import IMAGE_SLOT_NAME

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


@dataclass
class Issue:
    code: str
    severity: str            # error | warning
    slide: int               # 0-based, -1 = вся колода
    message: str
    bbox: list = field(default_factory=list)   # [x, y, w, h] в долях слайда
    id: str = ""             # стабильный идентификатор для UI и авто-фиксов
    deterministic: bool = True

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def at(cls, code, severity, slide, message, bbox=None, deterministic=True):
        return cls(code=code, severity=severity, slide=slide, message=message,
                   bbox=list(bbox or []), deterministic=deterministic)


class Audit:
    """Выполняет набор проверок над PDF/PPTX-артефактом."""

    def __init__(self, profile: dict):
        self.profile = profile
        allowed_hex = [p.get("hex") for p in profile.get("palette", []) if p.get("hex")]
        self.allowed_colors = set(c.upper() for c in allowed_hex)
        self.allowed_colors |= {c.upper() for c in profile.get("text_colors", [])}
        theme_colors = (profile.get("theme_colors") or {}).values()
        self.allowed_colors |= {c.upper() for c in theme_colors if c}
        self.fonts_allowed = set()
        for f in (profile.get("headline_font"), profile.get("body_font")):
            if f:
                self.fonts_allowed.add(f.lower())
        for fb in profile.get("fonts", []):
            self.fonts_allowed.add((fb.get("name") or "").lower())
        self.bg_color = self._default_bg()

        # разрешённая типографическая шкала и макеты шаблона
        scale = profile.get("type_scale", {}) or {}
        self.allowed_sizes = sorted({
            round(float(x), 2)
            for key in ("title", "body")
            for x in (scale.get(key) or [])
        })
        self.layouts = profile.get("layouts", []) or []
        self.layouts_by_name = {l.get("name"): l for l in self.layouts}
        self.layout_names = set(self.layouts_by_name)
        self.grid = profile.get("grid") or {}
        self.max_typefaces = int(profile.get("max_typefaces", 2))
        self.safe_area = self._compute_safe_area(profile)

    def _compute_safe_area(self, profile: dict) -> tuple[float, float, float, float]:
        """Безопасная зона слайда: ближайший к краю контент из всех макетов шаблона.

        Используется проверкой `content_in_margins`: если шаблон нигде не ставит
        контент ближе 1.2″ к краю, то и сгенерированный слайд не должен. Так
        проверка самосогласована: она не ругает вёрстку за то, что делает сам
        шаблон (у разных макетов поля разные).
        """
        size = profile.get("slide_size") or {}
        w = float(size.get("w_in") or 13.333)
        h = float(size.get("h_in") or 7.5)
        lefts, tops, rights, bottoms = [], [], [], []
        for layout in self.layouts:
            # и тело, и рамка заголовка: вертикальный заголовок шаблона scholar
            # стоит у самого края, и проверка полей не должна ругать вёрстку за
            # то, что делает сам макет
            boxes = [layout.get("body") or {}, layout.get("title_ph") or {}]
            for box in boxes:
                if box.get("w"):
                    lefts.append(float(box["x"]))
                    tops.append(float(box["y"]))
                    rights.append(w - (float(box["x"]) + float(box["w"])))
                    bottoms.append(h - (float(box["y"]) + float(box["h"])))
        grid = self.grid or {}
        if grid.get("body_w"):
            lefts.append(float(grid.get("body_x", 0.5)))
            tops.append(float(grid.get("body_y", 1.0)))
            rights.append(float(grid.get("margin_right", 0.5)))
            bottoms.append(float(grid.get("margin_bottom", 0.5)))
        if not lefts:
            return (0.2, 0.2, w - 0.4, h - 0.4)
        return (max(0.0, min(lefts)), max(0.0, min(tops)),
                w - max(0.0, min(rights)), h - max(0.0, min(bottoms)))

    def _default_bg(self) -> str:
        hs = [p.get("hex") for p in self.profile.get("palette", []) if p.get("hex")]
        for h in hs:
            if h and _lum(h) >= 0.6:
                return h.upper()
        return "#FFFFFF"

    # ------------------------------------------------------------------- main
    def audit(self, deck: Deck, pptx_bytes: bytes) -> dict:
        prs = Presentation(io.BytesIO(pptx_bytes))
        self.W, self.H = prs.slide_width, prs.slide_height
        # колода нужна проверкам, которым важен смысл слайда (тип, назначение):
        # витринные слайды вправе быть почти пустыми
        self._deck = deck
        issues: list[Issue] = []

        slide_objs = list(prs.slides)
        headings: dict[str, int] = {}
        fonts_per_slide: dict[int, set] = {}
        fingerprints: dict[str, int] = {}

        for si, slide in enumerate(slide_objs):
            geo = self._shape_geoms(slide)
            # 1. границы
            self._check_bounds(slide, geo, si, issues)
            # 2. наложения
            self._check_overlaps(slide, geo, si, issues)
            # 3. обрезка текста
            self._check_text_overflow(slide, si, issues)
            # 4. токены: шрифты, цвета, кегли
            self._check_tokens(slide, si, issues)
            self._check_font_sizes(slide, si, issues)
            self._check_fill_colors(slide, si, issues, geo)
            # 5. контраст: фактический фон (заливка фигуры → декор → p:bg макета)
            self._check_contrast(slide, si, issues)
            self._check_text_invisible(slide, si, issues)
            self._check_text_over_decor(slide, si, issues)
            self._check_empty_text_frames(slide, si, issues)
            # 6. плотность контента
            self._check_density(slide, si, issues)
            # 6b. объекты: таблицы, диаграммы, подписи, слоты изображений
            self._check_objects(slide, si, issues)
            # 6c. геометрия и композиция
            self._check_layout_from_template(slide, si, issues)
            self._check_content_area(slide, si, issues, geo)
            self._check_grid_alignment(slide, si, issues)
            self._check_images(slide, si, issues, geo)
            self._check_branding(slide, si, issues)
            self._check_fill_ratio(slide, si, issues, geo)
            # 7. пустой слайд
            self._check_empty(slide, si, issues, geo)
            # 8. заполнители-заглушки
            self._check_placeholders(slide, si, issues, geo)
            # 9. слайд-картинка
            self._check_raster_slide(slide, si, issues, geo)
            # заголовки и шрифты для сквозных проверок
            heading = self._heading(slide)
            if heading:
                headings[heading.lower()] = headings.get(heading.lower(), 0) + 1
            fonts_per_slide[si] = {it["font"].lower() for it in self._scan_runs(slide)
                                   if it["font"]}
            fingerprint = self._fingerprint(slide)
            if fingerprint:
                fingerprints[fingerprint] = fingerprints.get(fingerprint, 0) + 1

        for heading, n in headings.items():
            if n > 1:
                issues.append(Issue.at("duplicate_heading", "warning", -1,
                                       f"повторяющийся заголовок «{heading[:60]}» встречается {n} раза"))

        self._check_deck_typefaces(fonts_per_slide, issues)
        self._check_duplicate_slides(fingerprints, issues)

        for n, issue in enumerate(issues):
            issue.id = f"{issue.code}-{issue.slide}-{n}"

        n_err = sum(1 for i in issues if i.severity == "error")
        return {
            "passed": n_err == 0,
            "errors": n_err,
            "warnings": sum(1 for i in issues if i.severity == "warning"),
            "total": len(issues),
            "issues": [i.to_dict() for i in issues],
        }

    # ---------------------------------------------------------------- helpers
    @staticmethod
    def _fr(x, total):
        return 0.0 if total == 0 else max(0.0, min(1.0, x / total))

    def _shape_geoms(self, slide):
        out = []
        for sh in slide.shapes:
            if sh.left is None or sh.width is None:
                continue
            out.append({
                "shape": sh, "type": shape_type_int(sh),
                "x": sh.left, "y": sh.top, "w": sh.width, "h": sh.height,
                "filled": self._is_visible(sh),
            })
        return out

    @staticmethod
    def _is_visible(sh) -> bool:
        """Участвует ли фигура в анализе наложений (пустые рамки не считаем)."""
        st = shape_type_int(sh)
        if st == MSO_SHAPE_TYPE.GROUP.value:
            return False
        if st in (MSO_SHAPE_TYPE.TEXT_BOX.value, MSO_SHAPE_TYPE.PLACEHOLDER.value):
            try:
                return bool(sh.text_frame.text.strip())
            except Exception:  # noqa: BLE001
                return False
        return True

    def _heading(self, slide) -> str:
        try:
            if slide.shapes.title is not None:
                return slide.shapes.title.text_frame.text.strip()
        except Exception:
            pass
        return ""

    def _scan_runs(self, slide):
        """Все текстовые runs слайда с геометрией."""
        out = []
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            for p in tf.paragraphs:
                ptext = "".join(r.text for r in p.runs).strip()
                for r in p.runs:
                    if not r.text.strip():
                        continue
                    out.append({
                        "shape": sh, "run": r, "para": p, "ptext": ptext,
                        "font": self._run_font(r), "size": self._run_size(r),
                        "color": self._run_color(r), "bold": bool(r.font.bold) if r.font.bold is not None else False,
                    })
        return out

    @staticmethod
    def _run_font(r) -> str:
        try:
            n = r.font.name
            if n:
                return n
            rPr = r._r.rPr
            if rPr is not None:
                ea = rPr.find(f"{{{A_NS}}}ea")
                if ea is not None and ea.get("typeface"):
                    return ea.get("typeface")
                latin = rPr.find(f"{{{A_NS}}}latin")
                if latin is not None and latin.get("typeface"):
                    return latin.get("typeface")
        except Exception:
            pass
        return ""

    @staticmethod
    def _run_size(r) -> float:
        try:
            if r.font.size:
                return r.font.size.pt
        except Exception:
            pass
        return 0.0

    @staticmethod
    def _run_color(r) -> str:
        try:
            if r.font.color.type is not None:
                rgb = r.font.color.rgb
                if rgb is not None:
                    return str(rgb)
        except Exception:
            pass
        return ""

    # ---------------------------------------------------------------- checks
    def _check_bounds(self, slide, geo, si, issues):
        W, H = self.W, self.H
        for g in geo:
            x, y, w, h = g["x"], g["y"], g["w"], g["h"]
            if x < -10000 or y < -10000 or x + w > W + 10000 or y + h > H + 10000:
                issues.append(Issue.at(
                    "out_of_bounds", "error", si,
                    f"элемент «{g['shape'].name}» выходит за границы слайда",
                    [self._fr(x, W), self._fr(y, H), self._fr(w, W), self._fr(h, H)]))

    def _check_overlaps(self, slide, geo, si, issues):
        cur = []
        for g in geo:
            if not g["filled"] or g["w"] <= 0.02 * self.W:
                continue
            # если g лежит целиком внутри более крупной фигуры — это композиция карточки
            if self._contained_in(g, cur):
                continue
            for old in cur:
                if self._contained_in(g, [old]):
                    continue
                ox = max(g["x"], old["x"]); oy = max(g["y"], old["y"])
                ox2 = min(g["x"] + g["w"], old["x"] + old["w"])
                oy2 = min(g["y"] + g["h"], old["y"] + old["h"])
                iw, ih = ox2 - ox, oy2 - oy
                if iw <= 0 or ih <= 0:
                    continue
                if iw * ih > 0.18 * min(g["w"] * g["h"], old["w"] * old["h"]):
                    issues.append(Issue.at(
                        "overlap", "warning", si,
                        f"элементы «{old['shape'].name}» и «{g['shape'].name}» перекрываются",
                        [self._fr(ox, self.W), self._fr(oy, self.H),
                         self._fr(iw, self.W), self._fr(ih, self.H)]))
            cur.append(g)

    @staticmethod
    def _contained_in(g, boxes, tol=0.92):
        """Истина, если площадь g почти целиком внутри одной из заполненных рамок."""
        for b in boxes:
            if not b.get("filled"):
                continue
            ix = max(g["x"], b["x"]); iy = max(g["y"], b["y"])
            ix2 = min(g["x"] + g["w"], b["x"] + b["w"])
            iy2 = min(g["y"] + g["h"], b["y"] + b["h"])
            inter = max(0, ix2 - ix) * max(0, iy2 - iy)
            own = g["w"] * g["h"]
            if own and inter >= tol * own and b["w"] * b["h"] >= own:
                return True
        return False

    def _check_text_overflow(self, slide, si, issues):
        for sh in slide.shapes:
            if not sh.has_text_frame or sh.width is None:
                continue
            tf = sh.text_frame
            total = tf.text.strip()
            if not total:
                continue
            # суммарный кегль: максимум по runs; оценка — в дюймах
            estim = self._estimate_text_h(tf, sh.width)
            # допуск 10%: оценка высоты приблизительная, но заметное переполнение
            # (когда текст реально обрезается краем рамки) ловиться обязано.
            # ВАЖНО: sh.height — в EMU, поэтому сравниваем в одних единицах:
            # раньше дюймы сравнивались с EMU и проверка не срабатывала никогда.
            height_in = (sh.height or 0) / 914400.0
            if estim > 0 and height_in > 0 and estim > height_in * 1.10:
                issues.append(Issue.at(
                    "text_overflow", "warning", si,
                    f"текст «{total[:40]}…» может не поместиться в «{sh.name}»",
                    [self._fr(sh.left, self.W), self._fr(sh.top, self.H),
                     self._fr(sh.width, self.W), self._fr(sh.height, self.H)]))

    def _estimate_text_h(self, tf, w_emu) -> float:
        from ..layout.geometry import text_height_in
        total_h = 0.0
        for p in tf.paragraphs:
            psize = 0.0
            for r in p.runs:
                sz = self._run_size(r)
                if sz:
                    psize = max(psize, sz)
            if psize == 0:
                psize = 14.0
            txt = "".join(r.text for r in p.runs)
            w_in = max(0.1, float(w_emu) / 914400.0)
            total_h += text_height_in(txt, w_in, psize)
        return total_h  # inches

    def _check_tokens(self, slide, si, issues):
        for it in self._scan_runs(slide):
            fname = (it["font"] or "").lower()
            if fname and fname not in self.fonts_allowed:
                issues.append(Issue.at(
                    "font_not_allowed", "error", si,
                    f"шрифт «{it['font']}» не входит в дизайн-систему шаблона"))
            col = it["color"]
            if not col:
                continue
            colkey = col if col.startswith("#") else "#" + col
            colkey = colkey.upper()
            if colkey not in self.allowed_colors and colkey not in (self.allowed_extra()):
                issues.append(Issue.at(
                    "color_not_allowed", "error", si,
                    f"цвет «{colkey}» не входит в палитру шаблона",
                    self._shape_bbox(it["shape"], slide)))

    def allowed_extra(self):
        return {"#FFFFFF", "#000000"}

    def is_derived_color(self, hex_color: str, tolerance: float = 14.0) -> bool:
        """Является ли цвет смесью двух разрешённых цветов палитры.

        Вёрстка получает светлые подложки смешением акцента с фоном. Это
        осознанное производное токена, а не «левый» цвет, поэтому для ЗАЛИВОК
        такие оттенки разрешены (ADR-008). Для текста правило строже: только
        точные цвета палитры.
        """
        if not hex_color:
            return False
        try:
            target = hex_to_rgb(hex_color)
        except Exception:  # noqa: BLE001
            return False
        palette = [hex_to_rgb(h) for h in self.allowed_colors
                   if h and len(h.lstrip("#")) == 6]
        for first in palette:
            for second in palette:
                if self._point_on_segment(target, first, second, tolerance):
                    return True
        return False

    @staticmethod
    def _point_on_segment(point, start, end, tolerance: float) -> bool:
        """Лежит ли точка на отрезке start→end (с допуском по расстоянию)."""
        vector = [end[i] - start[i] for i in range(3)]
        length_sq = sum(v * v for v in vector)
        if length_sq == 0:
            return all(abs(point[i] - start[i]) <= tolerance for i in range(3))
        t = sum((point[i] - start[i]) * vector[i] for i in range(3)) / length_sq
        if t < -0.02 or t > 1.02:
            return False
        projection = [start[i] + t * vector[i] for i in range(3)]
        distance = sum((point[i] - projection[i]) ** 2 for i in range(3)) ** 0.5
        return distance <= tolerance

    def _check_fill_colors(self, slide, si, issues, geo):
        """Заливки должны быть из палитры шаблона или её смесями."""
        for g in geo:
            shape = g["shape"]
            if g["type"] not in (MSO_SHAPE_TYPE.AUTO_SHAPE.value,
                                 MSO_SHAPE_TYPE.TEXT_BOX.value):
                continue
            color = self._shape_fill(shape)
            if not color:
                continue
            key = "#" + color.lstrip("#").upper()[-6:]
            if key in self.allowed_colors or key in self.allowed_extra():
                continue
            if self.is_derived_color(key):
                continue
            issues.append(Issue.at(
                "color_not_allowed", "warning", si,
                f"заливка «{key}» не из палитры шаблона и не является её смесью",
                self._shape_bbox(shape, slide)))

    def _shape_bbox(self, sh, slide):
        if sh.left is None:
            return []
        return [self._fr(sh.left, self.W), self._fr(sh.top, self.H),
                self._fr(sh.width, self.W), self._fr(sh.height, self.H)]

    def _decor_fill_at(self, slide, shape) -> Optional[str]:
        """Заливка декора макета под фигурой: тот же расчёт, что в рендере.

        Декор наследуется от макета и не виден в фигурах слайда, но именно он
        определяет фактический фон текста: на тёмной плашке шаблона тёмный текст
        не читается, и проверять контраст к белому бессмысленно.
        """
        if shape.left is None or shape.width is None or shape.height is None:
            return None
        layout = self.layouts_by_name.get(self._layout_name(slide)) or {}
        x, y = shape.left / 914400, (shape.top or 0) / 914400
        w, h = shape.width / 914400, shape.height / 914400
        area = w * h
        if area <= 0:
            return None
        best_fill, best_ratio = None, 0.0
        for item in layout.get("decor", []) or []:
            fill = item.get("fill")
            if not fill:
                continue
            left = max(x, float(item.get("x", 0.0)))
            top = max(y, float(item.get("y", 0.0)))
            right = min(x + w, float(item.get("x", 0.0)) + float(item.get("w", 0.0)))
            bottom = min(y + h, float(item.get("y", 0.0)) + float(item.get("h", 0.0)))
            if right <= left or bottom <= top:
                continue
            ratio = ((right - left) * (bottom - top)) / area
            if ratio > best_ratio:
                best_fill, best_ratio = fill, ratio
        return best_fill if best_ratio >= 0.35 else None

    def _layout_bg(self, slide) -> Optional[str]:
        """Фактический фон макета из p:bg (ADR-034), иначе None."""
        layout = self.layouts_by_name.get(self._layout_name(slide)) or {}
        return layout.get("background")

    def _effective_bg(self, slide, shape) -> str:
        """Фон под текстом ровно так же, как его выбирает рендер.

        Приоритет: заливка самой фигуры → декор макета под рамкой → фон p:bg
        макета → самая светлая палитра. Единый алгоритм для рендера и аудита —
        иначе проверка контраста молчит на тёмных шаблонах.
        """
        return (self._shape_fill(shape)
                or self._decor_fill_at(slide, shape)
                or self._layout_bg(slide)
                or self.bg_color)

    def _check_contrast(self, slide, si, issues):
        for it in self._scan_runs(slide):
            col = it["color"]
            size = it["size"]
            if not col or not size:
                continue
            bg = self._effective_bg(slide, it["shape"])
            large = size >= 18 or (size >= 14 and it["bold"])
            threshold = 3.0 if large else 4.5
            r, g, b = hex_to_rgb(col)
            ratio = _contrast(r, g, b, *_hex(bg))
            # совсем невидимый текст (< 3:1) — отдельный код text_invisible;
            # contrast_too_low ловит «читается, но ниже нормы WCAG»
            if 3.0 <= ratio < threshold:
                req = f"{threshold:.1f}:1 (WCAG)"
                issues.append(Issue.at(
                    "contrast_too_low", "error", si,
                    f"контраст текста «#{col}» ниже {req}",
                    self._shape_bbox(it["shape"], slide)))

    def _check_text_invisible(self, slide, si, issues, floor: float = 3.0):
        """Текст и фон одного цвета: контраст ниже абсолютного порога 3:1.

        `contrast_too_low` использует полный порог WCAG (4.5:1 для мелкого),
        а эта проверка ловит именно «невидимый» текст — случай, когда цвет
        выбирался по палитре темы, а не по фактическому фону панели.
        """
        for it in self._scan_runs(slide):
            col = it["color"]
            if not col:
                continue
            bg = self._effective_bg(slide, it["shape"])
            r, g, b = hex_to_rgb(col)
            if _contrast(r, g, b, *_hex(bg)) < floor:
                issues.append(Issue.at(
                    "text_invisible", "error", si,
                    f"текст «{it['ptext'][:40]}» цвета «#{col}» не виден "
                    f"на фоне «{bg}» (контраст < {floor:.1f}:1)",
                    self._shape_bbox(it["shape"], slide)))

    def _check_empty_text_frames(self, slide, si, issues):
        """Пустая текстовая рамка рендера: текст потерялся при вёрстке.

        Проверяются только свободные текстовые рамки (не декор и не
        плейсхолдеры шаблона): таблицы, карточки-автофигуры и слоты картинок
        создаются без текста осознанно.
        """
        for sh in slide.shapes:
            try:
                if not sh.has_text_frame or sh.width is None or sh.height is None:
                    continue
                if shape_type_int(sh) != MSO_SHAPE_TYPE.TEXT_BOX.value:
                    continue
                if (sh.text_frame.text or "").strip():
                    continue
                area = (sh.width or 0) * (sh.height or 0)
                if area < 0.2 * 914400 * 914400:
                    continue
                if (sh.name or "").startswith(("TitleBox", IMAGE_SLOT_NAME)):
                    continue
                issues.append(Issue.at(
                    "empty_text_frame", "warning", si,
                    f"пустая текстовая рамка «{sh.name}»: текст потерялся",
                    self._shape_bbox(sh, slide)))
            except Exception:  # noqa: BLE001
                continue

    @staticmethod
    def _shape_fill(sh):
        """Заливка фигуры (solidFill srgbClr), иначе None."""
        try:
            if fill_type_int(sh) == MSO_FILL_TYPE.SOLID.value and sh.fill.fore_color is not None:
                return str(sh.fill.fore_color.rgb)
        except Exception:
            pass
        return None

    def _check_text_over_decor(self, slide, si, issues, min_cover: float = 0.12):
        """Текст не должен пересекать границу цветной панели декора.

        `_check_contrast` сравнивает буквы с доминирующим фоном под рамкой, но
        если рамка заходит на панель декора, часть текста ложится на другой фон.
        Белое на белом (или тёмное на тёмном) выглядит как обрывки текста —
        именно так терялись слайды на VK Education, при нулевых ошибках аудита.
        """
        layout = self.layouts_by_name.get(self._layout_name(slide)) or {}
        decor = layout.get("decor") or []
        # остаток рамки лежит на фоне слайда: сперва p:bg макета (честный фон
        # тёмных шаблонов), затем крупнейшая заливка слайда
        slide_bg = str(self._layout_bg(slide) or self._shape_bg(slide)).upper()
        for it in self._scan_runs(slide):
            shape = it["shape"]
            if shape.left is None or shape.width is None or shape.height is None:
                continue
            col, size = it["color"], it["size"]
            if not col or not size:
                continue
            large = size >= 18 or (size >= 14 and it["bold"])
            threshold = 3.0 if large else 4.5
            r, g, b = hex_to_rgb(col)
            own_fill = self._shape_fill(shape)
            if own_fill:
                # у фигуры своя непрозрачная заливка: декор под ней не виден
                regions = {str(own_fill).upper(): 1.0}
            elif not decor:
                continue
            else:
                x, y = shape.left / 914400, (shape.top or 0) / 914400
                w, h = shape.width / 914400, shape.height / 914400
                area = w * h
                if area <= 0:
                    continue
                regions = {}
                covered = 0.0
                for item in decor:
                    fill = item.get("fill")
                    if not fill:
                        continue
                    left = max(x, float(item.get("x", 0.0)))
                    top = max(y, float(item.get("y", 0.0)))
                    right = min(x + w, float(item.get("x", 0.0)) + float(item.get("w", 0.0)))
                    bottom = min(y + h, float(item.get("y", 0.0)) + float(item.get("h", 0.0)))
                    if right <= left or bottom <= top:
                        continue
                    ratio = (right - left) * (bottom - top) / area
                    key = str(fill).upper()
                    regions[key] = regions.get(key, 0.0) + ratio
                    covered += ratio
                # остаток рамки лежит на фоне слайда: там текст тоже обязан читаться.
                # Если в макете есть декор-картинки (полноэкранный фон), профиль
                # знает лишь средний цвет — остаток считаем тем же фоном, что и
                # доминирующий декор, иначе тёмные шаблоны дают ложное «белое на белом»
                remainder = max(0.0, 1.0 - min(1.0, covered))
                if remainder > 0:
                    rest_bg = slide_bg
                    if any(item.get("is_picture") for item in decor) and regions:
                        rest_bg = max(regions, key=regions.get)
                    regions[rest_bg] = regions.get(rest_bg, 0.0) + remainder
            # сначала проверяем доминирующий фон: если текст не читается уже на
            # нём, это обычный низкий контраст (contrast_too_low/text_invisible),
            # а не «край панели»
            dominant = max(regions, key=regions.get) if regions else ""
            if dominant and _contrast(r, g, b, *_hex(dominant)) < threshold:
                continue
            for key, ratio in regions.items():
                if ratio < min_cover:
                    continue
                if _contrast(r, g, b, *_hex(key)) < threshold:
                    issues.append(Issue.at(
                        "text_on_decor_edge", "warning", si,
                        f"часть текста «{it['ptext'][:40]}» лежит на фоне "
                        f"«{key}», где он не читается",
                        self._shape_bbox(shape, slide)))
                    break

    def _shape_bg(self, slide) -> str:
        fill_colors = []
        for sh in slide.shapes:
            try:
                if fill_type_int(sh) == MSO_FILL_TYPE.SOLID.value:
                    rgb = sh.fill.fore_color.rgb
                    if rgb is not None:
                        fill_colors.append((sh.width or 0) * (sh.height or 0), str(rgb))
            except Exception:
                continue
        if fill_colors:
            fill_colors.sort(reverse=True)
            return fill_colors[0][1]
        return self.bg_color

    def _check_density(self, slide, si, issues):
        runs = self._scan_runs(slide)
        # буллеты
        bullets = [it for it in runs if self._is_bullet(it["para"], it["run"])]
        unique_bullets = {it["ptext"] for it in bullets if it["ptext"]}
        if len(unique_bullets) > 6:
            issues.append(Issue.at("too_many_bullets", "error", si,
                                   f"на слайде {len(unique_bullets)} буллетов (максимум 6)"))
        for it in bullets:
            if len(it["ptext"].split()) > 15:
                issues.append(Issue.at("bullet_too_long", "error", si,
                                       f"буллет длиннее 15 слов: «{it['ptext'][:50]}…»",
                                       self._shape_bbox(it["shape"], slide)))

    def _check_objects(self, slide, si, issues):
        """Объекты данных: размеры таблиц, число серий, наличие картинки.

        Проверки таблиц и диаграмм разведены намеренно: раньше проверка размера
        таблицы была вложена в ветку диаграмм и не срабатывала никогда.
        """
        for sh in slide.shapes:
            try:
                if getattr(sh, "has_chart", False):
                    self._check_chart(sh, slide, si, issues)
                elif getattr(sh, "has_table", False):
                    self._check_table(sh, slide, si, issues)
                elif (sh.name or "").startswith(IMAGE_SLOT_NAME):
                    issues.append(Issue.at(
                        "image_missing", "warning", si,
                        "нет изображения для блока-иллюстрации: оставлен слот",
                        self._shape_bbox(sh, slide)))
            except Exception:  # noqa: BLE001 — не роняем аудит из-за одной фигуры
                continue

    def _check_table(self, sh, slide, si, issues):
        tbl = sh.table
        n_rows, n_cols = len(tbl.rows), len(tbl.columns)
        if n_rows > 7 or n_cols > 5:
            issues.append(Issue.at(
                "table_too_big", "error", si,
                f"таблица {n_rows}×{n_cols} (максимум 7 строк и 5 колонок)",
                self._shape_bbox(sh, slide)))

    def _check_chart(self, sh, slide, si, issues):
        chart = sh.chart
        n_series = sum(len(plot.series) for plot in chart.plots)
        if n_series > 5:
            issues.append(Issue.at(
                "too_many_series", "error", si,
                f"в диаграмме {n_series} серий (максимум 5)",
                self._shape_bbox(sh, slide)))
        self._check_chart_labels(sh, slide, si, issues, n_series)

    def _check_chart_labels(self, sh, slide, si, issues, n_series: int) -> None:
        """Диаграмма должна читаться без пояснений: легенда, категории, единицы.

        Проверяется: легенда при нескольких сериях, непустые категории и наличие
        подписей данных либо заголовка оси значений (единицы измерения).
        """
        chart = sh.chart
        problems: list[str] = []
        plots = list(chart.plots)
        if not plots:
            problems.append("нет данных")
        else:
            plot = plots[0]
            if n_series > 1 and not chart.has_legend:
                problems.append("нет легенды")
            if not list(plot.categories):
                problems.append("нет подписей категорий")
            has_labels = bool(getattr(plot, "has_data_labels", False))
            axis_titled = False
            if chart.chart_type not in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT):
                try:
                    axis_titled = bool(chart.value_axis.has_title)
                except Exception:  # noqa: BLE001
                    axis_titled = False
            if not has_labels and not axis_titled:
                problems.append("нет подписей данных и единиц измерения")
        if problems:
            issues.append(Issue.at(
                "chart_unlabeled", "warning", si,
                f"диаграмма нечитаема: {', '.join(problems)}",
                self._shape_bbox(sh, slide)))

    # ------------------------------------------------ шаблон, сетка, композиция
    def _check_font_sizes(self, slide, si, issues):
        """Кегль должен входить в типографическую шкалу шаблона."""
        if not self.allowed_sizes:
            return
        for it in self._scan_runs(slide):
            size = it["size"]
            if not size:
                continue
            if not any(abs(size - allowed) <= 0.6 for allowed in self.allowed_sizes):
                issues.append(Issue.at(
                    "font_size_not_in_scale", "error", si,
                    f"кегль {size:g} pt не из типографической шкалы шаблона "
                    f"({', '.join(f'{s:g}' for s in self.allowed_sizes[:8])}…)",
                    self._shape_bbox(it["shape"], slide)))

    def _check_layout_from_template(self, slide, si, issues):
        """Слайд обязан быть собран на макете из шаблона."""
        if not self.layout_names:
            return
        name = self._layout_name(slide)
        if name not in self.layout_names:
            issues.append(Issue.at(
                "layout_not_from_template", "error", si,
                f"слайд собран на макете «{name}», которого нет в шаблоне"))

    def _content_area(self, slide, si: int = -1) -> tuple[float, float, float, float] | None:
        """Контентная область слайда ровно так, как её выбирает вёрстка.

        Логика совпадает с `Renderer._canvas` (`layout.engine.choose_canvas`):
        тело макета берётся, только если оно достаточно велико, иначе
        используется сетка профиля; текст без собственной подложки не заходит на
        крупные панели декора. Иначе аудит сравнивал бы заполненность с
        «визиточной» рамкой, которую вёрстка осознанно не использует.
        """
        layout = self.layouts_by_name.get(self._layout_name(slide)) or {}
        if not layout:
            return None
        model = None
        if self._deck is not None and 0 <= si < len(self._deck.slides):
            model = self._deck.slides[si]
        rect = choose_canvas(self.profile, layout,
                             uniform_background=slide_needs_uniform_background(model))
        return (rect.x, rect.y, rect.w, rect.h)

    def _check_content_area(self, slide, si, issues, geo, tolerance_in: float = 0.08):
        """Контент не должен заходить в поля у краёв слайда.

        Поля берутся из безопасной зоны шаблона (см. `_compute_safe_area`), а не
        из одного макета: у разных макетов поля разные, и вёрстка имеет право
        использовать любую из них.
        """
        left, top, right, bottom = (v * 914400 for v in self.safe_area)
        tol = tolerance_in * 914400
        for shape in self._text_shapes(slide):
            if shape.left is None or shape.width is None:
                continue
            # заголовок стоит там, где его ставит вёрстка макета (обычно выше
            # контентной области) — проверка полей относится к контенту
            if self._is_title_shape(shape):
                continue
            if (shape.left < left - tol or (shape.top or 0) < top - tol
                    or shape.left + shape.width > right + tol
                    or (shape.top or 0) + (shape.height or 0) > bottom + tol):
                issues.append(Issue.at(
                    "content_in_margins", "warning", si,
                    f"блок «{shape.name}» заходит в поля у края слайда",
                    self._shape_bbox(shape, slide)))

    def _check_grid_alignment(self, slide, si, issues,
                              band: tuple[float, float] = (0.005, 0.09)):
        """Блоки одной колонки должны быть выровнены между собой.

        Сравниваются блоки, которые заведомо задуманы в одной колонке (близкие
        левые края и сопоставимая ширина). Микро-расхождение в пределах `band`
        (0.1–2.3 мм) выглядит как случайный сдвиг — это и есть «блок не выровнен
        по направляющим». Отступы внутренних элементов карточек больше band и
        не считаются ошибкой.
        """
        lo, hi = (v * 914400 for v in band)
        # блоки заголовочной зоны (заголовок и подзаголовок) в проверке не
        # участвуют: они выравниваются по плейсхолдеру заголовка макета
        header_limit = self.safe_area[1] * 914400 - 0.02 * 914400
        shapes = [s for s in self._text_shapes(slide)
                  if s.left is not None and (s.width or 0) > 0
                  and (s.top or 0) >= header_limit]
        for index, first in enumerate(shapes):
            for second in shapes[index + 1:]:
                delta = abs(first.left - second.left)
                if not (lo < delta <= hi):
                    continue
                # одна колонка — это не только близкий край, но и совпадающая
                # ширина: внутренние элементы блока (ячейки фактоидов, карточки)
                # заведомо уже родительской рамки
                if abs(first.width - second.width) > 0.01 * 914400:
                    continue
                issues.append(Issue.at(
                    "misaligned_to_grid", "warning", si,
                    f"блоки «{first.name}» и «{second.name}» не выровнены: "
                    f"расхождение левых краёв {delta / 914400:.3f}″",
                    self._shape_bbox(second, slide)))

    def _check_images(self, slide, si, issues, geo, tolerance: float = 0.05):
        """Пропорции картинки не должны быть нарушены (растяжение > 5%)."""
        for g in geo:
            if g["type"] != MSO_SHAPE_TYPE.PICTURE.value:
                continue
            shape = g["shape"]
            if not shape.width or not shape.height:
                continue
            try:
                native_w, native_h = shape.image.size
            except Exception:  # noqa: BLE001
                continue
            if not native_w or not native_h:
                continue
            source_ratio = native_w / native_h
            rendered_ratio = shape.width / shape.height
            if abs(rendered_ratio - source_ratio) / source_ratio > tolerance:
                issues.append(Issue.at(
                    "image_stretched", "error", si,
                    f"картинка «{shape.name}» растянута: пропорции "
                    f"{rendered_ratio:.2f} против исходных {source_ratio:.2f}",
                    self._shape_bbox(shape, slide)))

    def _check_branding(self, slide, si, issues, tolerance_in: float = 0.12):
        """Логотип и колонтитулы должны оставаться на местах макета.

        Проверяются только те фирменные элементы, которые реально перенесены на
        слайд: позиции берутся из профиля макета. Если шаблон держит логотип
        только на макете, проверка для слайда молчит — это ожидаемо, потому что
        наследуемый декор сдвинуть со слайда нельзя.
        """
        layout = self.layouts_by_name.get(self._layout_name(slide)) or {}
        branding = layout.get("branding") or []
        if not branding:
            return
        by_name = {}
        for shape in slide.shapes:
            try:
                by_name[shape.name] = shape
            except Exception:  # noqa: BLE001
                continue
        tol = tolerance_in * 914400
        for item in branding:
            shape = by_name.get(item.get("name"))
            if shape is None or shape.left is None:
                continue
            expected_x = float(item.get("x", 0.0)) * 914400
            expected_y = float(item.get("y", 0.0)) * 914400
            if abs(shape.left - expected_x) > tol or abs((shape.top or 0) - expected_y) > tol:
                issues.append(Issue.at(
                    "branding_shifted", "warning", si,
                    f"фирменный элемент «{item.get('name')}» "
                    f"({item.get('type')}) сдвинут с места макета",
                    self._shape_bbox(shape, slide)))

    CONTENT_TYPES = (MSO_SHAPE_TYPE.TABLE.value, MSO_SHAPE_TYPE.CHART.value,
                     MSO_SHAPE_TYPE.PICTURE.value, MSO_SHAPE_TYPE.AUTO_SHAPE.value,
                     MSO_SHAPE_TYPE.TEXT_BOX.value, MSO_SHAPE_TYPE.PLACEHOLDER.value)
    FILL_GRID = 100  # разрешение растеризации при подсчёте занятой площади

    def _check_fill_ratio(self, slide, si, issues, geo, sparse: float = 0.20,
                          dense: float = 1.10):
        """Заполненность контентной области: пусто или тесно.

        Считается доля площади, занятая контентными объектами (рамки текстовых
        блоков, таблицы, диаграммы, картинки) **относительно контентной области
        макета**, а не всего слайда. Причина: у части шаблонов контентная область
        по дизайну занимает треть слайда, и проверка «от слайда» ругала бы
        нормальную вёрстку (обоснование — docs/AUDIT.md).

        Заголовок, декор и полноэкранный фон не учитываются; витринные слайды
        (титул, раздел, финал) из проверки исключены — у них мало контента
        по замыслу макета.
        """
        layout = self.layouts_by_name.get(self._layout_name(slide)) or {}
        if layout.get("role") not in (None, "content", "agenda"):
            return
        if self._deck is not None and 0 <= si < len(self._deck.slides):
            # титул, раздел, оглавление и финал вправе быть лаконичными
            model = self._deck.slides[si]
            if model.slide_type != SlideType.CONTENT:
                return
            # слайд-фактоид (одна крупная цифра) и цитата лаконичны по замыслу:
            # заполненность рамок тут ничего не говорит о качестве вёрстки
            kinds = {b.kind for b in model.blocks}
            if kinds and kinds <= {"factoids", "quote"}:
                return
        area = self._content_area(slide, si)
        area_emu = ((area[2] * 914400) * (area[3] * 914400)) if area else (self.W * self.H)
        if area_emu <= 0:
            return
        grid = self.FILL_GRID
        cells: set[tuple[int, int]] = set()
        for g in geo:
            shape = g["shape"]
            if self._is_title_shape(shape) or not g["filled"]:
                continue
            if g["type"] not in self.CONTENT_TYPES:
                continue
            if g["type"] == MSO_SHAPE_TYPE.PICTURE.value \
                    and g["w"] >= 0.9 * self.W and g["h"] >= 0.9 * self.H:
                continue  # фон слайда, а не контент
            x0 = max(0, int(g["x"] / self.W * grid))
            x1 = min(grid, int((g["x"] + g["w"]) / self.W * grid) + 1)
            y0 = max(0, int(g["y"] / self.H * grid))
            y1 = min(grid, int((g["y"] + g["h"]) / self.H * grid) + 1)
            for cx in range(x0, x1):
                for cy in range(y0, y1):
                    cells.add((cx, cy))
        if not cells and not geo:
            return
        # доля занятых ячеек переводится в площадь и делится на контентную область
        covered = len(cells) / float(grid * grid) * (self.W * self.H)
        ratio = covered / area_emu
        if ratio < sparse:
            issues.append(Issue.at(
                "slide_too_sparse", "warning", si,
                f"контентная область заполнена на {ratio * 100:.0f}% "
                f"(меньше {sparse * 100:.0f}%)"))
        elif ratio > dense:
            issues.append(Issue.at(
                "slide_too_dense", "warning", si,
                f"контентная область переполнена: {ratio * 100:.0f}%"))

    @staticmethod
    def _layout_name(slide) -> str:
        try:
            return slide.slide_layout.name
        except Exception:  # noqa: BLE001
            return ""

    def _text_shapes(self, slide) -> list:
        """Уникальные текстовые фигуры слайда (без заголовка и без дублей по run)."""
        seen: set[int] = set()
        out = []
        for it in self._scan_runs(slide):
            shape = it["shape"]
            if id(shape) in seen or self._is_title_shape(shape):
                continue
            seen.add(id(shape))
            out.append(shape)
        return out

    @staticmethod
    def _is_title_shape(shape) -> bool:
        # макеты без рамки заголовка (шаблон scholar): вёрстка рисует заголовок
        # свободной рамкой с именем TitleBox — она тоже заголовок, а не контент
        if (shape.name or "").startswith("TitleBox"):
            return True
        try:
            if not shape.is_placeholder:
                return False
            return getattr(shape.placeholder_format.type, "value",
                           shape.placeholder_format.type) in (
                PP_PLACEHOLDER.TITLE.value, PP_PLACEHOLDER.CENTER_TITLE.value)
        except Exception:  # noqa: BLE001
            return False

    # ------------------------------------------------------------- сквозные
    def _fingerprint(self, slide) -> str:
        """Отпечаток содержимого слайда для поиска дублей."""
        parts = [self._heading(slide)]
        for it in self._scan_runs(slide):
            parts.append(it["ptext"])
        text = re.sub(r"\s+", " ", " ".join(parts)).strip().lower()
        return text if len(text) >= 24 else ""

    def _check_deck_typefaces(self, fonts_per_slide: dict, issues):
        """Гарнитур в колоде должно быть не больше двух (правило шаблона)."""
        counter: dict[str, int] = {}
        for fonts in fonts_per_slide.values():
            for font in fonts:
                counter[font] = counter.get(font, 0) + 1
        if len(counter) > self.max_typefaces:
            names = ", ".join(sorted(counter))
            issues.append(Issue.at(
                "too_many_typefaces", "warning", -1,
                f"в колоде {len(counter)} гарнитур ({names}); "
                f"рекомендуется не больше {self.max_typefaces}"))

    def _check_duplicate_slides(self, fingerprints: dict, issues):
        for fingerprint, count in fingerprints.items():
            if count > 1:
                issues.append(Issue.at(
                    "duplicate_slide", "warning", -1,
                    f"слайды дублируют друг друга: «{fingerprint[:60]}…» "
                    f"встречается {count} раза"))

    @staticmethod
    def _is_bullet(p, run) -> bool:
        try:
            pPr = p._pPr
            if pPr is None:
                return False
            # Clark-нотация обязательна: ElementTree ищет тег «{ns}имя», и без
            # фигурных скобок find() не находил buChar — проверки буллетов
            # (too_many_bullets, bullet_too_long) не срабатывали никогда
            if (pPr.find(f"{{{A_NS}}}buChar") is not None
                    or pPr.find(f"{{{A_NS}}}buAutoNum") is not None):
                return True
        except Exception:
            pass
        return False

    OBJECT_TYPES = (MSO_SHAPE_TYPE.CHART.value, MSO_SHAPE_TYPE.TABLE.value,
                    MSO_SHAPE_TYPE.PICTURE.value, MSO_SHAPE_TYPE.AUTO_SHAPE.value)

    def _check_empty(self, slide, si, issues, geo):
        has_text = any(it["run"].text.strip() for it in self._scan_runs(slide))
        has_obj = any(g["type"] in self.OBJECT_TYPES for g in geo)
        if not has_text and not has_obj and geo:
            issues.append(Issue.at("empty_slide", "warning", si,
                                   "слайд не содержит ни текста, ни объектов"))

    def _check_placeholders(self, slide, si, issues, geo):
        """Ищет служебную «рыбу» построчно, а не по всему тексту слайда.

        Построчная проверка не даёт ложных срабатываний на осмысленном тексте,
        где просто встретилось слово «заголовок».
        """
        for it in self._scan_runs(slide):
            text = (it["ptext"] or "").strip()
            if text and looks_like_placeholder(text):
                issues.append(Issue.at(
                    "placeholder_text", "warning", si,
                    f"текст-заглушка: «{text[:60]}»",
                    self._shape_bbox(it["shape"], slide)))
                return

    def _check_raster_slide(self, slide, si, issues, geo):
        W, H = self.W, self.H
        for g in geo:
            if g["type"] == MSO_SHAPE_TYPE.PICTURE.value \
                    and g["w"] >= 0.92 * W and g["h"] >= 0.92 * H:
                issues.append(Issue.at("raster_slide", "error", si,
                                       "слайд целиком является растровой картинкой",
                                       [0, 0, 1, 1]))


def fill_type_int(sh) -> int | None:
    """Числовой код типа заливки (см. `shape_type_int` — та же причина)."""
    fill = getattr(sh, "fill", None)
    ft = getattr(fill, "type", None)
    return getattr(ft, "value", ft)


def shape_type_int(sh) -> int | None:
    """Числовой код типа фигуры.

    `Shape.shape_type` возвращает член перечисления `MSO_SHAPE_TYPE`, который
    не сравнивается с int напрямую (`MSO_SHAPE_TYPE.PICTURE == 13` → False),
    поэтому везде приводим к значению.
    """
    st = getattr(sh, "shape_type", None)
    return getattr(st, "value", st)


def _lum(h: str) -> float:
    r, g, b = hex_to_rgb(h)
    return (0.2126 * r + 0.7152 * g + 0.0722 * b) / 255.0


def _hex(h: str) -> tuple:
    try:
        if len(h) == 8:
            h = h[2:]
        return hex_to_rgb(h)
    except Exception:
        return (0, 0, 0)


def _contrast(r1, g1, b1, r2, g2, b2) -> float:
    def L(r, g, b):
        rs = r / 255.0; gs = g / 255.0; bs = b / 255.0
        def f(c):
            return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
        return 0.2126 * f(rs) + 0.7152 * f(gs) + 0.0722 * f(bs)
    l1, l2 = L(r1, g1, b1), L(r2, g2, b2)
    if l1 < l2:
        l1, l2 = l2, l1
    return (l1 + 0.05) / (l2 + 0.05)