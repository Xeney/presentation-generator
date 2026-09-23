"""Импорт и декомпозиция контент-пакетов: PPTX, DOCX, TXT/MD → ContentCorpus.

Требование ТЗ: сервис должен импортировать и декомпозировать контент-пакеты, а не
принимать только текст брифа. PPTX разбирается по слайдам (заголовки, тезисы,
таблицы, диаграммы, изображения, цифры), DOCX и текст — по разделам (заголовок →
новый раздел), изображения DOCX вытаскиваются из `word/media`.
"""
from __future__ import annotations

import re
import uuid
from pathlib import Path
from zipfile import ZipFile

from .corpus import ContentCorpus, CorpusSlide, CorpusTable, find_numbers

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

IMAGE_EXT = {"png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp", "emf", "wmf"}
BULLET_PREFIX = re.compile(r"^\s*[-–—•*·]\s+")
MAX_IMAGES = 60


class ContentImportError(Exception):
    """Файл не удалось разобрать как контент-пакет."""


def import_content_pack(data: bytes, filename: str) -> ContentCorpus:
    """Разбирает файл контент-пакета, определяя формат по расширению."""
    suffix = Path(filename or "").suffix.lower()
    if suffix in (".pptx", ".pptm", ".potx"):
        corpus = _import_pptx(data, filename)
    elif suffix == ".docx":
        corpus = _import_docx(data, filename)
    elif suffix in (".txt", ".md", ".csv", ".json"):
        corpus = _import_text(data, filename)
    else:
        raise ContentImportError(
            f"формат {suffix or 'без расширения'} не поддерживается: "
            "ожидается PPTX, DOCX, TXT или MD")
    corpus.id = uuid.uuid4().hex[:12]
    corpus.source_file = filename
    corpus.numbers = corpus.all_numbers()
    if not corpus.non_empty_slides():
        raise ContentImportError("в файле не найдено текстового содержимого")
    return corpus


# --------------------------------------------------------------------- PPTX
def _iter_shapes(shapes):
    """Рекурсивно обходит фигуры, разворачивая группы."""
    for shape in shapes:
        try:
            if shape.shape_type is not None and str(shape.shape_type).startswith("GROUP"):
                yield from _iter_shapes(shape.shapes)
                continue
        except Exception:  # noqa: BLE001
            pass
        yield shape


def _shape_text(shape) -> list[tuple[str, int]]:
    """Пары (текст абзаца, уровень отступа) для текстовой фигуры."""
    out: list[tuple[str, int]] = []
    if not getattr(shape, "has_text_frame", False):
        return out
    for paragraph in shape.text_frame.paragraphs:
        text = re.sub(r"\s+", " ", paragraph.text or "").strip()
        if text:
            out.append((text, int(paragraph.level or 0)))
    return out


def _is_title(shape) -> bool:
    try:
        return shape.is_placeholder and int(shape.placeholder_format.type) in (1, 3)
    except Exception:  # noqa: BLE001
        return False


def _is_subtitle(shape) -> bool:
    try:
        return shape.is_placeholder and int(shape.placeholder_format.type) == 4
    except Exception:  # noqa: BLE001
        return False


def _table_data(shape) -> CorpusTable | None:
    if not getattr(shape, "has_table", False):
        return None
    try:
        table = shape.table
        rows = [[re.sub(r"\s+", " ", cell.text or "").strip() for cell in row.cells]
                for row in table.rows]
    except Exception:  # noqa: BLE001
        return None
    if not rows:
        return None
    return CorpusTable(header=rows[0], rows=rows[1:])


def _chart_data(shape) -> dict | None:
    if not getattr(shape, "has_chart", False):
        return None
    try:
        chart = shape.chart
        plots = list(chart.plots)
        categories = [str(c) for c in plots[0].categories] if plots else []
        series = []
        for plot in plots:
            for s in plot.series:
                series.append({"name": str(s.name), "values": [float(v) for v in s.values]})
        return {"type": str(chart.chart_type), "categories": categories, "series": series}
    except Exception:  # noqa: BLE001
        return None


def _picture_bytes(shape) -> tuple[str, bytes] | None:
    """Байты картинки фигуры (в том числе picture-плейсхолдера), иначе None."""
    try:
        image = shape.image
    except Exception:  # noqa: BLE001 — у не-картинок атрибута нет
        return None
    ext = (image.ext or "png").lower()
    if ext not in IMAGE_EXT:
        return None
    return ext, image.blob


def _import_pptx(data: bytes, filename: str) -> ContentCorpus:
    import io

    from pptx import Presentation

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ContentImportError(f"PPTX не открывается: {exc}") from exc

    corpus = ContentCorpus(kind="pptx")
    for index, slide in enumerate(prs.slides):
        item = CorpusSlide(index=index, layout=slide.slide_layout.name)
        texts: list[str] = []
        for shape in _iter_shapes(slide.shapes):
            if _is_title(shape):
                parts = _shape_text(shape)
                if parts and not item.heading:
                    item.heading = parts[0][0]
                    texts.append(parts[0][0])
                continue
            if _is_subtitle(shape):
                parts = _shape_text(shape)
                if parts and not item.subheading:
                    item.subheading = parts[0][0]
                    texts.append(parts[0][0])
                continue
            table = _table_data(shape)
            if table is not None:
                item.tables.append(table)
                texts.append(" ".join(table.header))
                texts.extend(" ".join(row) for row in table.rows)
                continue
            chart = _chart_data(shape)
            if chart is not None:
                item.charts.append(chart)
                texts.append(" ".join(chart["categories"]))
                texts.extend(" ".join(str(v) for v in s["values"]) for s in chart["series"])
                continue
            picture = _picture_bytes(shape)
            if picture is not None and len(corpus.images) < MAX_IMAGES:
                ext, blob = picture
                key = f"s{index + 1}_img{len(item.images) + 1}.{ext}"
                corpus.images[key] = blob
                corpus.image_labels[key] = (
                    f"слайд {index + 1} «{item.heading or slide.slide_layout.name}», "
                    f"фигура «{shape.name}»")
                item.images.append(key)
                continue
            for text, level in _shape_text(shape):
                texts.append(text)
                is_list_shape = _is_body(shape)
                if level > 0 or is_list_shape or BULLET_PREFIX.match(text):
                    item.bullets.append(BULLET_PREFIX.sub("", text))
                else:
                    item.paragraphs.append(text)

        try:
            if slide.has_notes_slide:
                notes = re.sub(r"\s+", " ", slide.notes_slide.notes_text_frame.text or "").strip()
                item.notes = notes[:600]
        except Exception:  # noqa: BLE001
            item.notes = ""

        item.bullets = _dedupe(item.bullets)[:12]
        item.paragraphs = _dedupe(item.paragraphs)[:8]
        item.numbers = find_numbers(" ".join(texts))
        corpus.slides.append(item)
    return corpus


def _is_body(shape) -> bool:
    try:
        return shape.is_placeholder and int(shape.placeholder_format.type) in (2, 7)
    except Exception:  # noqa: BLE001
        return False


def _dedupe(items: list[str]) -> list[str]:
    seen: dict[str, None] = {}
    for item in items:
        seen.setdefault(item, None)
    return list(seen)


# --------------------------------------------------------------------- DOCX
def _import_docx(data: bytes, filename: str) -> ContentCorpus:
    from lxml import etree

    try:
        with ZipFile(__import__("io").BytesIO(data)) as zf:
            names = zf.namelist()
            if "word/document.xml" not in names:
                raise ContentImportError("в DOCX нет word/document.xml")
            root = etree.fromstring(zf.read("word/document.xml"))
            media = [(n, zf.read(n)) for n in names if n.startswith("word/media/")]
    except ContentImportError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ContentImportError(f"DOCX не открывается: {exc}") from exc

    corpus = ContentCorpus(kind="docx")
    current = CorpusSlide(index=0)
    texts: list[str] = []
    table_buffer: list[CorpusTable] = []

    def flush():
        nonlocal current, texts, table_buffer
        current.tables = table_buffer
        current.numbers = find_numbers(" ".join(texts))
        corpus.slides.append(current)
        current = CorpusSlide(index=len(corpus.slides))
        texts, table_buffer = [], []

    def add_text(text: str, style: str = "", is_list: bool = False):
        nonlocal current
        text = re.sub(r"\s+", " ", text or "").strip()
        if not text:
            return
        texts.append(text)
        lowered = (style or "").lower()
        if "heading" in lowered or "заголов" in lowered:
            if current.heading or current.bullets or current.paragraphs or current.tables:
                flush()
                corpus.slides[-1].index = len(corpus.slides) - 1
                current.index = len(corpus.slides) - 1
            current.heading = text
        elif is_list or BULLET_PREFIX.match(text):
            current.bullets.append(BULLET_PREFIX.sub("", text))
        else:
            current.paragraphs.append(text)

    for node in root.iter():
        tag = etree.QName(node).localname
        if tag == "p":
            text = "".join(t.text or "" for t in node.iter(f"{{{W_NS}}}t"))
            style = ""
            pstyle = node.find(f"{{{W_NS}}}pPr/{{{W_NS}}}pStyle")
            if pstyle is not None:
                style = pstyle.get(f"{{{W_NS}}}val", "")
            is_list = node.find(f"{{{W_NS}}}pPr/{{{W_NS}}}numPr") is not None
            add_text(text, style, is_list)
        elif tag == "tbl":
            rows = []
            for tr in node.findall(f"{{{W_NS}}}tr"):
                cells = []
                for tc in tr.findall(f"{{{W_NS}}}tc"):
                    cells.append(" ".join(
                        re.sub(r"\s+", " ", t.text or "").strip()
                        for t in tc.iter(f"{{{W_NS}}}t")).strip())
                if any(cells):
                    rows.append(cells)
            if rows:
                table_buffer.append(CorpusTable(header=rows[0], rows=rows[1:]))
                texts.extend(" ".join(row) for row in rows)

    flush()
    if corpus.slides and corpus.slides[-1].is_empty():
        corpus.slides.pop()

    for n, blob in media[:MAX_IMAGES]:
        ext = Path(n).suffix.lstrip(".").lower()
        if ext not in IMAGE_EXT:
            continue
        key = f"doc_img{len(corpus.images) + 1}.{ext}"
        corpus.images[key] = blob
        corpus.image_labels[key] = f"изображение из DOCX ({Path(n).name})"
    if corpus.slides and corpus.images:
        corpus.slides[0].images = list(corpus.images)
    return corpus


# ---------------------------------------------------------------- TXT / MD
def _import_text(data: bytes, filename: str) -> ContentCorpus:
    for encoding in ("utf-8", "cp1251", "utf-16"):
        try:
            raw = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ContentImportError("не удалось определить кодировку текста")

    corpus = ContentCorpus(kind="text")
    current = CorpusSlide(index=0)
    for raw_line in raw.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("#") or (line.startswith("==") and line.endswith("==")):
            if current.heading or current.bullets or current.paragraphs or current.tables:
                corpus.slides.append(current)
                current = CorpusSlide(index=len(corpus.slides))
            current.heading = line.lstrip("#= ").rstrip("= ").strip()
        elif BULLET_PREFIX.match(line):
            current.bullets.append(BULLET_PREFIX.sub("", line))
        elif "|" in line and line.count("|") >= 2:
            cells = [c.strip() for c in line.strip("|").split("|")]
            if set("".join(cells)) <= set("-: "):
                continue
            if not current.tables:
                current.tables.append(CorpusTable(header=cells))
            else:
                current.tables[0].rows.append(cells)
        elif line.endswith(":") and len(line) < 80:
            current.subheading = line.rstrip(":")
        else:
            current.paragraphs.append(line)
    corpus.slides.append(current)
    for slide in corpus.slides:
        slide.numbers = find_numbers(" ".join(
            slide.bullets + slide.paragraphs + [slide.heading, slide.subheading]))
    return corpus
