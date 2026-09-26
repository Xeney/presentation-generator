"""HTML+CSS → PPTX нативными объектами python-pptx (ADR-036).

Разбирает разметку `html_renderer` (section.slide + абсолютно позиционированные
`.el`) и собирает из неё нативные объекты: текстовые фреймы с runs, списки с
буллетами и автонумерацией, таблицы, диаграммы, шевроны, автофигуры подложек.
Растеризация запрещена: слайд целиком картинкой не становится, `<img>` попадает
в PPTX только как контентная картинка внутри своего слота.
"""
from __future__ import annotations

import io
import json
import logging
import re
from typing import Any, Optional

from lxml import html as lxml_html
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .pptx_renderer import Renderer

log = logging.getLogger("html_to_pptx")

PCT = re.compile(r"(-?[0-9.]+)%")
PT = re.compile(r"(-?[0-9.]+)pt", re.I)
PX = re.compile(r"(-?[0-9.]+)px", re.I)
HEX = re.compile(r"#([0-9a-fA-F]{6})")
RGB = re.compile(r"rgba?\(\s*([0-9]+)\s*,\s*([0-9]+)\s*,\s*([0-9]+)")
VAR = re.compile(r"var\((--[a-z0-9-]+)(?:\s*,\s*([^)]+))?\)", re.I)

CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "donut": XL_CHART_TYPE.DOUGHNUT,
}


class HtmlToPptxError(RuntimeError):
    pass


# ------------------------------------------------------------------ CSS-утилиты
def _parse_root_vars(html: str) -> dict[str, str]:
    match = re.search(r":root\s*\{(.*?)\}", html, re.S)
    if not match:
        return {}
    out: dict[str, str] = {}
    for piece in match.group(1).split(";"):
        if ":" not in piece:
            continue
        key, value = piece.split(":", 1)
        key = key.strip()
        if key.startswith("--"):
            out[key] = value.strip()
    return out


def _style_dict(element) -> dict[str, str]:
    style = element.get("style") or ""
    out: dict[str, str] = {}
    for piece in style.split(";"):
        if ":" not in piece:
            continue
        key, value = piece.split(":", 1)
        out[key.strip().lower()] = value.strip()
    return out


def _resolve(value: str, variables: dict[str, str], depth: int = 0) -> str:
    if not value or depth > 5:
        return value

    def replace(match: re.Match) -> str:
        name, fallback = match.group(1), match.group(2)
        return variables.get(name, (fallback or "").strip())

    return _resolve(VAR.sub(replace, value), variables, depth + 1)


def _color(value: str | None, variables: dict[str, str]) -> Optional[RGBColor]:
    if not value:
        return None
    value = _resolve(value, variables)
    match = HEX.search(value) or RGB.search(value)
    if not match:
        return None
    if match.re is HEX:
        return RGBColor.from_string(match.group(1).upper())
    r, g, b = (int(match.group(i)) for i in (1, 2, 3))
    return RGBColor(r, g, b)


def _pt(value: str | None, variables: dict[str, str], default: Optional[float] = None,
        snap=None) -> Optional[float]:
    if not value:
        return default
    value = _resolve(value, variables)
    match = PT.search(value) or PX.search(value)
    size = float(match.group(1)) if match else default
    if size is None:
        return None
    if snap is not None:
        try:
            size = snap(size)
        except Exception:  # noqa: BLE001
            pass
    return float(size)


def _font_name(value: str | None, variables: dict[str, str], default: str) -> str:
    if not value:
        return default
    value = _resolve(value, variables)
    first = value.split(",")[0].strip().strip("'\"")
    return first or default


def _emu_from_pct(value: str | None, total_in: float, variables: dict[str, str]) -> int:
    if not value:
        return 0
    value = _resolve(value, variables)
    match = PCT.search(value)
    if not match:
        return 0
    return int(float(match.group(1)) / 100.0 * total_in * 914400)


def _emu_from_inches(value: float) -> int:
    return int(value * 914400)


def _one_hex(color) -> str:
    return f"#{color}" if color is not None else ""


# ------------------------------------------------------------------- конвертер
class HtmlToPptxConverter:
    """DOM из `html_renderer` → PPTX нативными объектами."""

    def __init__(self, profile: dict, template_bytes: bytes | None = None):
        self.profile = profile
        self.template_bytes = template_bytes
        size = profile.get("slide_size") or {}
        self.slide_w = float(size.get("w_in", 13.333))
        self.slide_h = float(size.get("h_in", 7.5))
        self.renderer = Renderer(profile, variant="compact",
                                 template_bytes=template_bytes)
        self.default_font = profile.get("body_font") or "Arial"
        self.headline_font = profile.get("headline_font") or self.default_font

    # ---------------------------------------------------------------- публичное
    def convert(self, html: str) -> bytes:
        document = lxml_html.fromstring(html)
        variables = _parse_root_vars(html)
        sections = document.xpath("//section[contains(@class, 'slide')]")
        if not sections:
            raise HtmlToPptxError("в HTML нет ни одного слайда (section.slide)")
        if self.template_bytes:
            prs = Presentation(io.BytesIO(self.template_bytes))
            Renderer._drop_original_slides(prs)
        else:
            prs = Presentation()
            prs.slide_width = Emu(_emu_from_inches(self.slide_w))
            prs.slide_height = Emu(_emu_from_inches(self.slide_h))
        for section in sections:
            self._add_slide(prs, section, variables)
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------ слайд
    def _add_slide(self, prs: Presentation, section, variables: dict[str, str]) -> None:
        layout_id = section.get("data-layout") or ""
        layout = None
        for profile_layout in self.profile.get("layouts", []):
            if profile_layout.get("id") == layout_id:
                layout = profile_layout
                break
        layout_obj = (self.renderer._layout_object(prs, layout)
                      if layout else prs.slide_masters[0].slide_layouts[0])
        slide = prs.slides.add_slide(layout_obj)
        self._drop_unused_placeholders(slide)

        title = section.xpath("./h1[contains(@class, 'title')]")
        title_element = title[0] if title else None
        if title_element is not None:
            self._draw_title(slide, title_element, variables)

        for element in section.xpath("./*"):
            tag = element.tag.lower() if isinstance(element.tag, str) else ""
            if tag in ("h1", "nav", "script", "style"):
                continue
            try:
                self._draw_element(slide, element, variables)
            except Exception as exc:  # noqa: BLE001 — один узел не должен ронять слайд
                log.warning("не удалось конвертировать узел %s: %s", tag, exc)

    @staticmethod
    def _drop_unused_placeholders(slide) -> None:
        keep = {1, 3}  # title, center title
        for placeholder in list(slide.placeholders):
            try:
                ph_type = getattr(placeholder.placeholder_format.type, "value",
                                  placeholder.placeholder_format.type)
                if ph_type in keep:
                    continue
                if placeholder.has_text_frame and not placeholder.text_frame.text.strip():
                    placeholder._element.getparent().remove(placeholder._element)
            except Exception:  # noqa: BLE001
                continue

    def _draw_title(self, slide, element, variables: dict[str, str]) -> None:
        text = "".join(element.itertext()).strip()
        if not text:
            return
        style = _style_dict(element)
        color = _color(style.get("color"), variables) or RGBColor(0x1A, 0x1A, 0x1A)
        size = _pt(style.get("font-size"), variables, 24.0,
                   snap=self.renderer._snap_size)
        font = _font_name(style.get("font-family"), variables, self.headline_font)
        title_placeholder = None
        try:
            title_placeholder = slide.shapes.title
        except Exception:  # noqa: BLE001
            title_placeholder = None
        if title_placeholder is not None:
            frame = title_placeholder.text_frame
            frame.clear()
            run = frame.paragraphs[0].add_run()
            run.text = text
            _apply_font(run, font, size, True, color)
            return
        box = slide.shapes.add_textbox(
            Emu(_emu_from_pct(style.get("left"), self.slide_w, variables)),
            Emu(_emu_from_pct(style.get("top"), self.slide_h, variables)),
            Emu(_emu_from_pct(style.get("width"), self.slide_w, variables)),
            Emu(_emu_from_pct(style.get("height"), self.slide_h, variables)))
        box.name = "TitleBox"
        frame = box.text_frame
        frame.word_wrap = True
        run = frame.paragraphs[0].add_run()
        run.text = text
        _apply_font(run, font, size, True, color)

    # ---------------------------------------------------------------- элементы
    def _draw_element(self, slide, element, variables: dict[str, str]) -> None:
        tag = element.tag.lower() if isinstance(element.tag, str) else ""
        classes = set((element.get("class") or "").split())
        style = _style_dict(element)
        rect = self._rect(style, variables)
        if tag == "ul":
            self._text_list(slide, element, rect, style, variables, numbered=False)
            return
        if tag == "ol":
            self._text_list(slide, element, rect, style, variables, numbered=True)
            return
        if tag == "table":
            self._table(slide, element, rect, style, variables)
            return
        if tag == "div" and "chart" in classes:
            self._chart(slide, element, rect, style, variables)
            return
        if tag == "img" and "image" in classes:
            self._image(slide, element, rect)
            return
        if tag == "blockquote":
            self._textbox(slide, element, rect, style, variables, italic=True, prefix="«",
                          suffix="»")
            return
        if "chevron" in classes:
            self._chevron(slide, element, rect, style, variables)
            return
        if "card" in classes or "bar" in classes or "factoid-card" in classes \
                or "image-slot" in classes:
            self._shape(slide, element, rect, style, variables)
            return
        self._textbox(slide, element, rect, style, variables)

    def _rect(self, style: dict, variables: dict[str, str]) -> dict[str, int]:
        return {
            "x": _emu_from_pct(style.get("left"), self.slide_w, variables),
            "y": _emu_from_pct(style.get("top"), self.slide_h, variables),
            "w": _emu_from_pct(style.get("width"), self.slide_w, variables),
            "h": _emu_from_pct(style.get("height"), self.slide_h, variables),
        }

    def _textbox(self, slide, element, rect: dict, style: dict,
                 variables: dict[str, str], *, italic: bool = False,
                 prefix: str = "", suffix: str = "") -> None:
        text = "".join(element.itertext()).strip()
        if not text:
            return
        if prefix or suffix:
            text = f"{prefix}{text}{suffix}"
        color = _color(style.get("color"), variables) or RGBColor(0, 0, 0)
        size = _pt(style.get("font-size"), variables, 14.0,
                   snap=self.renderer._snap_size)
        font = _font_name(style.get("font-family"), variables, self.default_font)
        weight = style.get("font-weight") or ""
        underline = "underline" in (style.get("text-decoration") or "")
        box = slide.shapes.add_textbox(
            Emu(rect["x"]), Emu(rect["y"]), Emu(rect["w"]), Emu(rect["h"]))
        frame = box.text_frame
        frame.word_wrap = True
        run = frame.paragraphs[0].add_run()
        run.text = text
        run.font.underline = underline if underline else None
        _apply_font(run, font, size, weight.strip() in ("700", "bold"), color,
                    italic=italic)

    def _text_list(self, slide, element, rect: dict, style: dict,
                   variables: dict[str, str], *, numbered: bool) -> None:
        items = ["".join(item.itertext()).strip()
                 for item in element.xpath("./li")]
        items = [item for item in items if item]
        if not items:
            return
        color = _color(style.get("color"), variables) or RGBColor(0, 0, 0)
        bullet_color = _color(style.get("--bullet-color"), variables) or color
        size = _pt(style.get("font-size"), variables, 14.0,
                   snap=self.renderer._snap_size)
        font = _font_name(style.get("font-family"), variables, self.default_font)
        box = slide.shapes.add_textbox(
            Emu(rect["x"]), Emu(rect["y"]), Emu(rect["w"]), Emu(rect["h"]))
        frame = box.text_frame
        frame.word_wrap = True
        for index, item in enumerate(items):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = PP_ALIGN.LEFT
            _set_list_marker(paragraph, bullet_color, numbered)
            run = paragraph.add_run()
            run.text = item
            _apply_font(run, font, size, False, color)

    def _table(self, slide, element, rect: dict, style: dict,
               variables: dict[str, str]) -> None:
        rows = element.xpath("./tbody/tr") or element.xpath("./tr")
        headers = element.xpath("./thead/tr/th") or element.xpath("./tr[1]/th")
        if not rows and not headers:
            return
        header_texts = ["".join(cell.itertext()).strip() for cell in headers]
        n_cols = max([len(header_texts)] + [len(row.xpath("./td")) for row in rows]) or 1
        n_rows = len(rows) + (1 if header_texts else 0)
        if n_cols == 0 or n_rows == 0:
            return
        header_bg = _color(style.get("--header-bg"), variables)
        header_fg = _color(style.get("--header-color"), variables) or RGBColor(255, 255, 255)
        cell_fg = _color(style.get("--cell-color"), variables) or RGBColor(0, 0, 0)
        zebra_bg = _color(style.get("--zebra-bg"), variables)
        font = _font_name(style.get("font-family"), variables, self.default_font)
        graphic = slide.shapes.add_table(n_rows, n_cols, Emu(rect["x"]), Emu(rect["y"]),
                                         Emu(rect["w"]), Emu(rect["h"]))
        table = graphic.table
        try:
            table.first_row = False
            table.horz_banding = False
            tbl_el = graphic._element.graphic.graphicData.tbl
            tblPr = tbl_el.tblPr
            tblPr.set("firstRow", "0")
            tblPr.set("bandRow", "0")
        except Exception:  # noqa: BLE001
            pass
        size = 11.0
        for col, text in enumerate(header_texts[:n_cols]):
            cell = table.cell(0, col)
            _fill_cell(cell, header_bg, text, font, size, header_fg, bold=True)
        offset = 1 if header_texts else 0
        for row_index, row in enumerate(rows):
            cells = row.xpath("./td")
            for col in range(n_cols):
                text = "".join(cells[col].itertext()).strip() if col < len(cells) else ""
                cell = table.cell(row_index + offset, col)
                fill = zebra_bg if (row_index % 2 == 1) else None
                _fill_cell(cell, fill, text, font, size, cell_fg, bold=False)

    def _chart(self, slide, element, rect: dict, style: dict,
               variables: dict[str, str]) -> None:
        payload_raw = element.get("data-chart") or ""
        if not payload_raw:
            return
        try:
            payload = json.loads(payload_raw)
        except json.JSONDecodeError:
            return
        chart_type = CHART_TYPES.get(str(payload.get("type", "column")).lower(),
                                     XL_CHART_TYPE.COLUMN_CLUSTERED)
        data = CategoryChartData()
        data.categories = payload.get("categories") or []
        for series in payload.get("series") or []:
            data.add_series(series.get("name") or "Ряд", series.get("values") or [])
        graphic = slide.shapes.add_chart(chart_type, Emu(rect["x"]), Emu(rect["y"]),
                                         Emu(rect["w"]), Emu(rect["h"]), data)
        chart = graphic.chart
        is_pie = chart_type in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT)
        chart.has_legend = (len(payload.get("series") or []) > 1) and not is_pie
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        try:
            plot = chart.plots[0]
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.number_format_is_linked = False
            labels.number_format = '0"%"' if payload.get("unit") == "%" else "General"
            if is_pie:
                labels.show_percentage = True
                labels.show_value = False
        except Exception:  # noqa: BLE001
            pass
        colors = payload.get("colors") or []
        try:
            for index, series in enumerate(list(chart.plots[0].series)):
                color = colors[index % len(colors)] if colors else None
                if color:
                    _set_series_color(series._element, str(color))
        except Exception:  # noqa: BLE001
            pass
        _style_chart_text(chart, self.renderer)

    def _image(self, slide, element, rect: dict) -> None:
        src = element.get("src") or ""
        match = re.match(r"data:(image/[a-z]+);base64,(.+)", src, re.S)
        if not match:
            return
        import base64

        try:
            blob = base64.b64decode(match.group(2))
        except Exception:  # noqa: BLE001
            return
        box = _contain_box(rect, blob)
        picture = slide.shapes.add_picture(io.BytesIO(blob), Emu(box["x"]),
                                           Emu(box["y"]), Emu(box["w"]),
                                           Emu(box["h"]))
        picture.name = "image"

    def _chevron(self, slide, element, rect: dict, style: dict,
                 variables: dict[str, str]) -> None:
        fill = _color(style.get("background-color"), variables)
        color = _color(style.get("color"), variables) or RGBColor(255, 255, 255)
        size = _pt(style.get("font-size"), variables, 12.0,
                   snap=self.renderer._snap_size)
        font = _font_name(style.get("font-family"), variables, self.default_font)
        shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Emu(rect["x"]), Emu(rect["y"]),
                                       Emu(rect["w"]), Emu(rect["h"]))
        shape.adjustments[0] = 0.25
        if fill is not None:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
        shape.shadow.inherit = False
        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        text = "".join(element.itertext()).strip()
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = text
        _apply_font(run, font, size, True, color)

    def _shape(self, slide, element, rect: dict, style: dict,
               variables: dict[str, str]) -> None:
        fill = _color(style.get("background-color"), variables)
        line = _color(style.get("border-color"), variables)
        classes = set((element.get("class") or "").split())
        rounded = "card" in classes or "factoid-card" in classes or "image-slot" in classes
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(shape_type, Emu(rect["x"]), Emu(rect["y"]),
                                       Emu(rect["w"]), Emu(rect["h"]))
        if "image-slot" in classes:
            # имя слота: аудит должен видеть отсутствие картинки и в HTML-пути
            from .pptx_renderer import IMAGE_SLOT_NAME

            shape.name = IMAGE_SLOT_NAME
        if rounded:
            try:
                shape.adjustments[0] = 0.06
            except Exception:  # noqa: BLE001
                pass
        if fill is not None:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill
        else:
            shape.fill.background()
        border = style.get("border") or ""
        border_color = _color(border, variables) or line
        if border_color is not None:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False
        text = "".join(element.itertext()).strip()
        if text:
            color = _color(style.get("color"), variables) or RGBColor(0, 0, 0)
            size = _pt(style.get("font-size"), variables, 12.0,
                       snap=self.renderer._snap_size)
            font = _font_name(style.get("font-family"), variables, self.default_font)
            frame = shape.text_frame
            frame.word_wrap = True
            run = frame.paragraphs[0].add_run()
            run.text = text
            _apply_font(run, font, size, False, color)


# --------------------------------------------------------------- низкоуровневое
def _apply_font(run, name: str, size: Optional[float], bold: bool, color: RGBColor,
                italic: bool = False) -> None:
    run.font.bold = bold
    run.font.italic = italic
    if size:
        run.font.size = Pt(float(size))
    if color is not None:
        run.font.color.rgb = color
    if name:
        run.font.name = name
        rPr = run._r.get_or_add_rPr()
        rPr.set("lang", "ru-RU")
        for tag in ("latin", "ea", "cs"):
            element = rPr.find(qn(f"a:{tag}"))
            if element is None:
                element = rPr.makeelement(qn(f"a:{tag}"), {})
                rPr.append(element)
            element.set("typeface", name)


def _set_list_marker(paragraph, color: RGBColor, numbered: bool) -> None:
    pPr = paragraph._pPr if paragraph._pPr is not None else paragraph._p.get_or_add_pPr()
    pPr.set("marL", "228600" if numbered else "182880")
    pPr.set("indent", "-228600" if numbered else "-182880")
    for tag in ("buNone", "buChar", "buAutoNum", "buClr", "buFont"):
        element = pPr.find(qn(f"a:{tag}"))
        if element is not None:
            pPr.remove(element)
    buClr = pPr.makeelement(qn("a:buClr"), {})
    buClr.append(pPr.makeelement(qn("a:srgbClr"), {"val": str(color)}))
    pPr.append(buClr)
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    if numbered:
        pPr.append(pPr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"}))
    else:
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "▪"}))


def _fill_cell(cell, fill: Optional[RGBColor], text: str, font: str, size: float,
               color: RGBColor, *, bold: bool) -> None:
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame = cell.text_frame
    frame.word_wrap = True
    frame.clear()
    run = frame.paragraphs[0].add_run()
    run.text = text
    _apply_font(run, font, size, bold, color)
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.background()


def _contain_box(rect: dict, blob: bytes) -> dict:
    """Вписывает картинку в прямоугольник с сохранением пропорций (contain).

    CSS `object-fit: contain` в PPTX не существует: если вставить картинку по
    размеру рамки, она растянется (аудит: `image_stretched`). Считаем ту же
    пропорциональную вставку, что делает python-pptx-рендерер.
    """
    try:
        from PIL import Image

        with Image.open(io.BytesIO(blob)) as image:
            width, height = image.size
    except Exception:  # noqa: BLE001
        return dict(rect)
    if not width or not height:
        return dict(rect)
    ratio = width / height
    box_w = float(rect["w"])
    box_h = box_w / ratio
    if box_h > rect["h"]:
        box_h = float(rect["h"])
        box_w = box_h * ratio
    return {
        "x": int(rect["x"] + (rect["w"] - box_w) / 2),
        "y": int(rect["y"] + (rect["h"] - box_h) / 2),
        "w": int(box_w),
        "h": int(box_h),
    }


def _set_series_color(series_element, hex_color: str) -> None:
    spPr = series_element.find(qn("c:spPr"))
    if spPr is None:
        return
    solid = spPr.find(qn("a:solidFill"))
    if solid is not None:
        spPr.remove(solid)
    solid = spPr.makeelement(qn("a:solidFill"), {})
    srgb = spPr.makeelement(qn("a:srgbClr"), {"val": hex_color.lstrip("#")})
    solid.append(srgb)
    spPr.append(solid)


def _style_chart_text(chart, renderer: Renderer) -> None:
    color = renderer._text_color()
    size = renderer._chart_font_size()
    font = renderer._body_font()
    targets: list[Any] = []
    try:
        if chart.has_legend:
            targets.append(chart.legend)
    except Exception:  # noqa: BLE001
        pass
    try:
        targets.append(chart.value_axis.tick_labels)
        targets.append(chart.category_axis.tick_labels)
    except Exception:  # noqa: BLE001
        pass
    for target in targets:
        try:
            target.font.size = Pt(size)
            target.font.name = font
            target.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        except Exception:  # noqa: BLE001
            continue


def html_to_pptx(html: str, profile: dict,
                 template_bytes: bytes | None = None) -> bytes:
    """HTML колоды → PPTX нативными объектами (без растеризации)."""
    return HtmlToPptxConverter(profile, template_bytes=template_bytes).convert(html)
