"""Проверка видимости текста в готовом PPTX: контраст к фактическому фону.

Фактический фон определяется тем же алгоритмом, что в рендере и аудите:
заливка фигуры → декор макета → фон p:bg макета → палитра. Полезно для
сквозной проверки на тёмных шаблонах без открытия картинок.

    python tools/check_visibility.py deck.pptx                 # профиль из шаблона рядом
    python tools/check_visibility.py deck.pptx template.pptx
"""
from __future__ import annotations

import argparse
import io
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from pptx import Presentation  # noqa: E402

from app.audit.checks import Audit, _contrast, _hex  # noqa: E402
from app.render.images import hex_to_rgb  # noqa: E402
from app.template.parser import TemplateParser  # noqa: E402


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="Контраст текста к фактическому фону")
    parser.add_argument("deck", type=Path)
    parser.add_argument("template", nargs="?", type=Path, default=None)
    parser.add_argument("--quiet", action="store_true")
    args = parser.parse_args(argv)

    template = args.template
    if template is None:
        candidate = args.deck.with_name("template.pptx")
        if candidate.exists():
            template = candidate
    if template is None or not template.exists():
        print("нужен шаблон: укажите вторым аргументом или положите template.pptx рядом")
        return 2

    profile = TemplateParser(template.read_bytes()).parse().to_dict()
    prs = Presentation(io.BytesIO(args.deck.read_bytes()))
    auditor = Audit(profile)
    auditor.W, auditor.H = prs.slide_width, prs.slide_height

    invisible = below = total = 0
    for si, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            bg = auditor._effective_bg(slide, shape)
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    if not run.text.strip():
                        continue
                    try:
                        color = "#" + str(run.font.color.rgb)
                    except Exception:  # noqa: BLE001
                        continue
                    size = run.font.size.pt if run.font.size else 14.0
                    bold = bool(run.font.bold)
                    threshold = 3.0 if (size >= 18 or (size >= 14 and bold)) else 4.5
                    ratio = _contrast(*hex_to_rgb(color), *_hex(bg))
                    total += 1
                    if ratio < 3.0:
                        invisible += 1
                        print(f"  слайд {si + 1}: «{run.text[:40]}» {color} на {bg}: "
                              f"контраст {ratio:.2f}")
                    elif ratio < threshold:
                        below += 1
    print(f"{args.deck.name}: текстов {total}, невидимых (<3:1) {invisible}, "
          f"ниже WCAG {below}")
    return 1 if invisible else 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
